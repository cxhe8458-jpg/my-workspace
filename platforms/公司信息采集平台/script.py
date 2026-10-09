# -*- coding: utf-8 -*-
"""
================================================================================
LLM 驱动的公司基础信息采集 Agent（DeepSeek）—— 通用周期化脚本
================================================================================
每次运行都会在 runs/ 下创建独立目录 runs/<时间戳_批次名>/，各周期互不混杂：
  runs/<tag>/
    输入名单/                 本次名单文件留档（溯源）
    companies_info.xlsx      最终产物（25 字段）
    logos/                   下载的 logo 图片
    html/                    抓到的首页/关于页/联系页 HTML
    .cache/companies_info.json  续跑缓存，每批落盘，崩溃后靠它恢复
    processed_domains.txt    已抓域名清单（断点续跑）
    failed_urls.txt          失败清单
    company_agent.log        运行日志

流程：
  [1] 传入公司网址（名单文件按“官网/网址”列读取，自动跳过无官网/非网址值）
      —— 同时读取“公司名称”列作为【原公司名称】写入结果，便于比对名单是否写错
  [2] 多线程、多方式获取首页 HTML（requests 桌面/移动 UA / 万网CDN解码 / Playwright 渲染）
  [3] 首页 HTML 落盘： html/<公司名称>.html
  [2b] 从首页提取并下载 logo 图片： logos/<域名>.<ext>
  [4] AI 解析首页 -> 公司名称、产品大类、“关于公司”栏目 URL
  [5] 抓“关于公司”页 HTML -> html/关于<公司名称>信息.html
  [7] AI 解析关于页 -> 25 字段 Schema（见 PROMPT_STEP7）
  [7b] 抓“联系我们”页并解析，补全邮箱/电话/传真/地址
  [8] 汇总到 .cache/companies_info.json（每批落盘，续跑靠它恢复）
  [9] 导出 companies_info.xlsx —— 唯一正式产物

字段（25 个；中文与英文内容各占一列不合并，*为必填）：
  原公司名称 original_name     名单里登记的原始名称（比对名单是否有误）
  *中文名称 cn_name            *英文名称 name（域名首字母大写）   *官网地址 website
  中文地址 cn_address          英文地址 address
  联系电话 mobile              固定电话 phone                  传真机 fax
  邮箱 email                   *属性 attribute（制造商/代理商/学校研究所）
  *国别 nation（ISO 3166-1 两位编码）   工厂国别 factory_nation
  别名 alias                   中文概述 cn_overview             英文概述 overview
  口号 slogan                  标签 tags                       行业 industry
  主营产品 main_products       *产品一级分类 category_level1   *产品二级分类 category_level2
  log图片 logo（自动下载）      信息来源 data_source            厂家性质 company_nature

分类名单：工作目录下放 产品分类表.csv（两列：一级分类,二级分类）或 分类名单.xlsx /
product_categories.json 时，产品一/二级分类会严格按名单映射；没有则用内置参考词表。

logo 图片：自动下载并统一转换为 PNG 格式（Pillow 转换，SVG 用 Playwright 渲染截图）。

输入名单（优先级从高到低）：
  urls.txt（每行 url 或 url,公司名称）
  --list 指定的名单文件
  PS_厂家积累表.xlsx（按“官网”列读取）
  原始网址名单/原始网址_待采集.xlsx + 剩余名单.xlsx（按域名合并去重）
  公司名单_已清洗.xlsx / 公司名单.xlsx（旧单列格式，兼容保留）

运行：
  python script.py                          # 新周期：自动建 runs/<时间戳>/ 并跑全流程
  python script.py --run-name 光博会2026夏  # 运行目录名加备注
  python script.py --list 新名单.xlsx       # 指定本次名单文件
  python script.py --resume <tag>           # 续跑已有周期目录（断点续跑）
  python script.py --export none            # 只抓取、不导出
  python script.py --export-only --resume <tag>   # 仅把该周期缓存导出 Excel
  python refill.py --tag <tag> [域名...]    # 对某周期缺字段记录定向补采

依赖：pip install requests beautifulsoup4 openpyxl   （可选：playwright 渲染兜底）
================================================================================
"""

import os
import re
import gc
import csv
import json
import time
import shutil
import logging
import argparse
import threading
import datetime
import traceback
from collections import OrderedDict
from urllib.parse import urlparse, urljoin
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from requests.adapters import HTTPAdapter
try:
    from urllib3.util.retry import Retry
except Exception:
    from requests.packages.urllib3.util.retry import Retry
from bs4 import BeautifulSoup
import urllib3
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

try:
    from playwright.sync_api import sync_playwright
    _HAS_PLAYWRIGHT = True
except Exception:
    _HAS_PLAYWRIGHT = False


# ==============================================================================
# ★★★ 一键运行配置（新增公司时，只需修改下面这一处）★★★
#   用法：把 URL 改成新公司的官网网址，然后直接运行   python script.py
#   其余所有配置、逻辑、输出格式均无需改动；URL 留空则走原有名单流程。
# ==============================================================================
URL = ""          # ← 【唯一需要修改的位置】例：URL = "https://www.qtlaser.cn/"
# ==============================================================================


# ==============================================================================
# 一、配置
# ==============================================================================
# DeepSeek（OpenAI 兼容接口）。建议用环境变量覆盖密钥：set DEEPSEEK_API_KEY=...
DEEPSEEK_API_KEY = os.environ.get("DEEPSEEK_API_KEY", "sk-YOUR_API_KEY_HERE")
DEEPSEEK_API_URL = "https://api.deepseek.com/chat/completions"
DEEPSEEK_MODEL = "deepseek-v4-flash"     # 可用模型：deepseek-v4-flash / deepseek-v4-pro
# 关掉思维链。deepseek-v4-flash 是推理模型，思维链(reasoning_content)**计入 max_tokens**，
# 而本项目全是"照着 HTML 填固定字段"的抽取任务，长篇推理毫无收益却会把额度烧光
# （详见 deepseek_chat 的注释）。实测关掉后：快约 6 倍、零烧穿、字段逐一对照完全一致。
AI_REASONING_EFFORT = "none"             # 置 None 可恢复推理模式
AI_MAX_TOKENS = 8192                     # 不再被思维链挤占，纯给正文

# ---- 当前生效的 AI 接口参数 ----
# 上面几个常量只是**默认值**。真正生效的以 server_config.json 的 ai 段为准，
# 那是网页「设置」页写的——换模型、换供应商随时改，不用动代码、不用重启服务。
# 配置文件不存在（命令行单跑 script.py）时就用上面的常量，互不依赖。
AI_CONFIG_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "server_config.json")
_ai_cfg_cache = {"mtime": None, "cfg": None}


def ai_config():
    """返回当前生效的 AI 接口参数 dict。按文件 mtime 缓存，不会每次调用都读盘。"""
    cfg = {
        "api_url": DEEPSEEK_API_URL,
        "api_key": DEEPSEEK_API_KEY,
        "model": DEEPSEEK_MODEL,
        "max_tokens": AI_MAX_TOKENS,
        "temperature": 0,
        "timeout": 180,
        "reasoning_effort": AI_REASONING_EFFORT or "",
        "json_mode": True,
    }
    try:
        mtime = os.path.getmtime(AI_CONFIG_FILE)
    except OSError:
        return cfg
    if _ai_cfg_cache["mtime"] == mtime and _ai_cfg_cache["cfg"]:
        return dict(_ai_cfg_cache["cfg"])
    try:
        with open(AI_CONFIG_FILE, "r", encoding="utf-8") as f:
            ai = (json.load(f) or {}).get("ai") or {}
    except Exception:
        ai = {}                       # 配置坏了不能让抓取起不来，回落默认
    for k in ("api_url", "api_key", "model"):
        if ai.get(k):
            cfg[k] = str(ai[k]).strip()
    for k in ("max_tokens", "timeout"):
        try:
            if ai.get(k):
                cfg[k] = int(ai[k])
        except (TypeError, ValueError):
            pass
    if ai.get("temperature") is not None:
        try:
            cfg["temperature"] = float(ai["temperature"])
        except (TypeError, ValueError):
            pass
    # 这两个允许被显式置空/置假，所以判"键在不在"而不是判真值
    if "reasoning_effort" in ai:
        cfg["reasoning_effort"] = str(ai.get("reasoning_effort") or "").strip()
    if "json_mode" in ai:
        cfg["json_mode"] = bool(ai["json_mode"])
    _ai_cfg_cache.update(mtime=mtime, cfg=dict(cfg))
    return cfg

# ---- 周期目录：每次运行一个独立文件夹 runs/<时间戳_批次名>/，互不混杂 ----
RUNS_DIR = "runs"
INPUTS_DIR = "输入名单"                    # 运行目录内：名单文件留档子目录
RUN_TAG = None                             # 当前运行周期标识（init_run_env 时赋值）
RUN_DIR = None                             # 当前运行周期目录（init_run_env 时赋值）

# 以下路径为“相对运行周期目录”的初始值，init_run_env() 里会切换为实际路径
HTML_DIR = "html"                          # 步骤3/6：HTML 落盘目录
LOGO_DIR = "logos"                         # 步骤2b：logo 图片目录
EXPORT_DIR = "."                           # 步骤9：导出目录（运行周期目录）
OUTPUT_XLSX = "companies_info.xlsx"        # 步骤9：最终产物
DEFAULT_EXPORT = "excel"                   # excel | none（csv/both 已废弃）
CACHE_DIR = ".cache"                       # 续跑缓存目录（每批落盘，崩溃后靠它恢复）
OUTPUT_JSON = os.path.join(CACHE_DIR, "companies_info.json")
PROCESSED_FILE = "processed_domains.txt"   # 已处理域名清单（断点续跑，防重复抓取）
FAILED_FILE = "failed_urls.txt"            # 失败清单（便于后续重试）
LOG_FILE = "company_agent.log"             # 运行日志（init_run_env 后落进运行目录）

# 分类名单候选文件（存在则严格按名单映射产品一/二级分类）
# CSV 格式：两列，第一列一级分类、第二列二级分类（如 产品分类表.csv）
CATEGORY_FILES = ["产品分类表.csv", "分类名单.xlsx", "产品分类.xlsx", "categories.xlsx",
                  "产品分类.json", "product_categories.json"]

FETCH_WORKERS = 8                         # 抓取线程
AI_WORKERS = 4                            # AI 解析线程（按接口并发限制调小）
CHUNK_SIZE = 100                          # 分批大小：每批跑完即落盘 JSON 并回收内存（大批量抓取必备）
TIMEOUT = 25
MAX_RETRIES = 3
VERIFY_SSL = False
USE_PLAYWRIGHT = True
PW_WAIT_UNTIL = "domcontentloaded"        # Playwright 等待策略；不用 networkidle(长连接/万网站永不idle会超时)
PW_RENDER_WAIT = 5000                     # DOM 就绪后再等待 JS 渲染的毫秒数(SPA/万网建站)
PW_NAV_TIMEOUT = 45                       # Playwright 单页导航超时(秒)，比 requests 宽松
HTML_MAX_CHARS = 40000                    # 送入模型前的 HTML 截断上限（控费/防超长）
MIN_VALID_HTML = 200                      # 正文小于此值视为抓取失败/需换方式
META_DESC_MIN = 80                        # meta description 达到此长度时计入有效正文(服务端渲染但可见正文短的站)

# 待处理公司网址（也可放到 urls.txt，每行一个，存在则优先读取）
URLS = [
  "http://www.gem-oe.com/"
]

# 25 个抓取字段：(内部 key, Excel 表头)。带 * 的为需求中的必填字段。
# 中文/英文内容各占一列，不合并（地址、概述均拆为两个字段）。
SCHEMA_FIELDS = [
    ("original_name",   "原公司名称"),    # 名单里登记的原始名称，便于比对名单是否有误
    ("cn_name",         "中文名称"),      # *1
    ("name",            "英文名称"),      # *2
    ("website",         "官网地址"),      # *3
    ("cn_address",      "中文地址"),      # 4
    ("address",         "英文地址"),      # 4
    ("mobile",          "联系电话"),      # 5
    ("phone",           "固定电话"),      # 6
    ("fax",             "传真机"),        # 7
    ("email",           "邮箱"),          # 8
    ("attribute",       "属性"),          # *9
    ("nation",          "国别"),          # *10
    ("factory_nation",  "工厂国别"),      # 11
    ("alias",           "别名"),          # 12
    ("cn_overview",     "中文概述"),      # 13
    ("overview",        "英文概述"),      # 13
    ("slogan",          "口号"),          # 14
    ("tags",            "标签"),          # 15
    ("industry",        "行业"),          # 16
    ("main_products",   "主营产品"),      # 17
    ("category_level1", "产品一级分类"),  # *18
    ("category_level2", "产品二级分类"),  # *19
    ("logo",            "log图片"),       # 20
    ("data_source",     "信息来源"),      # 21
    ("company_nature",  "厂家性质"),      # 22
]
SCHEMA_KEYS = [k for k, _h in SCHEMA_FIELDS]     # 兼容旧引用（内部 key 顺序）


# ------------------------------------------------------------------------------
# 提示词
# ------------------------------------------------------------------------------
PROMPT_STEP4_SYS = ("你是一位从业十年的资深爬虫工程师，知晓所有网页结构，能解析所有 HTML 文件。")
PROMPT_STEP4_USER = (
    "请解析下面这段 HTML，提取：公司名称、产品名称（产品大类名称，数组）、"
    "“关于我们/关于公司/公司概况”栏目信息；"
    "最重要的是找出两个栏目对应的 URL（href，可为相对路径）："
    "①“关于我们/关于公司/公司概况” 的 URL；②“联系我们/联系方式” 的 URL。\n"
    "若某个栏目不存在则该字段返回空字符串。只输出 JSON，格式如下（不要额外文字）：\n"
    '{"cn_name":"公司中文名","products":["大类1","大类2"],'
    '"about_url":"关于栏目的href","contact_url":"联系我们栏目的href","about_summary":"关于栏目简述"}\n'
    "HTML 内容如下：\n```html\n%s\n```"
)

# 在“关于我们”页里再次找“联系我们”URL（首页没找到时的兜底）
PROMPT_FIND_CONTACT_SYS = "你是一位资深爬虫工程师，擅长从 HTML 链接中定位栏目入口。"
PROMPT_FIND_CONTACT_USER = (
    "请从下面 HTML 的链接中找出“联系我们/联系方式”栏目对应的 URL（href，可为相对路径）。"
    "只输出 JSON：{\"contact_url\":\"href或空字符串\"}\n"
    "HTML 内容如下：\n```html\n%s\n```"
)

# 联系页解析：提取联系类字段（字段名与主 Schema 一致，便于合并）
PROMPT_CONTACT_SYS = ("你是一个从业十年的 HTML 解析专家，请获取里面的公司信息包括电话、邮箱、传真、地址等信息。")
PROMPT_CONTACT_USER = (
    "请解析下面的“联系我们”页 HTML，提取联系信息，只输出 JSON（不要额外文字），"
    "字段与要求如下（无则置 null；中英文地址分开输出，各自只放一种语言）：\n"
    '{\n'
    '  "email": "邮箱地址（只要一个有效邮箱）",\n'
    '  "mobile": "移动电话（+86-1xxxxxxxxxx 格式）",\n'
    '  "phone": "固定电话（按公司所在地加国家区号，如 +86-xxx-xxxxxx）",\n'
    '  "fax": "传真（+86-xxx-xxxxxx）",\n'
    '  "cn_address": "公司中文地址（官网没有中文地址时翻译成中文）",\n'
    '  "address": "公司英文地址（官网没有英文地址时翻译成英文）"\n'
    '}\n'
    "HTML 内容如下：\n```html\n%s\n```"
)

# 步骤7：按 25 字段需求解析关于页（中英文内容拆分为独立字段）
PROMPT_STEP7_SYS = "角色: 你是一位从业15年的资深爬虫工程师，知晓所有网页结构，能解析所有html文件"
PROMPT_STEP7_USER = (
    "请解析下面的 HTML，按以下字段输出一个 JSON 对象（只输出 JSON，不要额外文字，无信息则置 null）。\n"
    "字段含义与要求（中文与英文分成两个独立字段，各自只放一种语言）：\n"
    '{\n'
    '  "cn_name": "中文名称(必填)：完整有效的中文名称，例如 xxx有限责任公司；只有英文名称的厂家不要自动翻译，直接写官网上的英文名称",\n'
    '  "name": "英文名称：留空即可（程序会用域名自动生成）",\n'
    '  "website": "官网地址：留空即可（程序自动填写）",\n'
    '  "cn_address": "中文地址：公司中文地址；官网只有英文地址时翻译成中文",\n'
    '  "address": "英文地址：公司英文地址；官网只有中文地址时翻译成英文",\n'
    '  "mobile": "联系电话：移动号码，+86-1xxxxxxxxxx 类型格式",\n'
    '  "phone": "固定电话：固定号码，+86-xxx-xxxxxx 格式",\n'
    '  "fax": "传真机：+86-xxx-xxxxxx 格式",\n'
    '  "email": "邮箱：一个有效邮箱地址",\n'
    '  "attribute": "属性(必填)：厂家属于制造商、代理商、学校研究所（可多选，用;分隔，主要属性放前面）",\n'
    '  "nation": "国别(必填)：厂家主体公司国别（ISO 3166-1 两位英文编码，如 CN/US/JP/DE/GB/KR；国外公司在中国的分公司，国别填国外主体公司所在国）",\n'
    '  "factory_nation": "工厂国别：国外公司在国内有生产基地的填 CN；国内公司的生产基地填 CN；其余与国别一致",\n'
    '  "alias": "别名：曾用名、别名、简称等，多个用;分隔",\n'
    '  "cn_overview": "中文概述：官网上的厂家简介（中文，通顺组合控制在200字左右；官网只有英文简介时翻译成中文）",\n'
    '  "overview": "英文概述：厂家简介的英文版（官网只有中文简介时翻译成英文）",\n'
    '  "slogan": "口号：厂家口号，简短，控制字符数",\n'
    '  "tags": "标签：厂家在相应产品类型、厂家类型等方面的短标签，多个用;分隔",\n'
    '  "industry": "行业：厂家生产产品所处相关行业标签，多个用;分隔",\n'
    '  "main_products": "主营产品：主要产品内容/产品名称，多个用;分隔",\n'
    '  "category_level1": "产品一级分类(必填)：对厂家的产品内容映射当前的产品一级分类；包含多类一级产品的厂家选取产品最多的前两个，用;分隔",\n'
    '  "category_level2": "产品二级分类(必填)：对厂家的产品内容映射当前的产品二级分类；包含多个二级产品的厂家选取产品最多的前两个，用;分隔",\n'
    '  "company_nature": "厂家性质：按厂家主营产品和产品分类等信息综合判断，从[激光器、相机、光学元件、光谱仪、其它]中选择一个",\n'
    '  "data_source": "信息来源：填 official_website",\n'
    '  "logo": "log图片：填 null（程序自动下载）"\n'
    '}\n'
    "已知该公司网址为：%s\n"
    "首页解析出的产品大类（供分类参考）：%s\n"
    "%s\n"
    "HTML 内容如下：\n```html\n%s\n```"
)


# ==============================================================================
# 二、日志 & 会话
# ==============================================================================
def setup_logger():
    logger = logging.getLogger("company_agent")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    fmt = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s", datefmt="%H:%M:%S")
    ch = logging.StreamHandler(); ch.setFormatter(fmt); logger.addHandler(ch)
    return logger


log = setup_logger()


def init_run_env(run_name=None, resume_tag=None):
    """初始化本次运行周期目录 runs/<tag>/（新周期用时间戳命名，可加批次备注；
    续跑时传入 --resume 的 tag 复用原目录），并把所有产物路径切到该目录。"""
    global RUN_DIR, RUN_TAG, HTML_DIR, LOGO_DIR, EXPORT_DIR, CACHE_DIR, OUTPUT_JSON
    global PROCESSED_FILE, FAILED_FILE, LOG_FILE
    if resume_tag:
        tag = str(resume_tag).strip().strip("/\\")
        if tag.startswith(RUNS_DIR + os.sep):
            tag = tag[len(RUNS_DIR) + 1:]
    else:
        stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        suffix = sanitize_filename(run_name) if run_name else ""
        tag = stamp + ("_" + suffix if suffix else "")
    run_dir = os.path.abspath(os.path.join(RUNS_DIR, tag))
    RUN_TAG, RUN_DIR = tag, run_dir
    HTML_DIR = os.path.join(run_dir, "html")
    LOGO_DIR = os.path.join(run_dir, "logos")
    EXPORT_DIR = run_dir
    CACHE_DIR = os.path.join(run_dir, ".cache")
    OUTPUT_JSON = os.path.join(CACHE_DIR, "companies_info.json")
    PROCESSED_FILE = os.path.join(run_dir, "processed_domains.txt")
    FAILED_FILE = os.path.join(run_dir, "failed_urls.txt")
    LOG_FILE = os.path.join(run_dir, "company_agent.log")
    for d in (run_dir, HTML_DIR, LOGO_DIR, CACHE_DIR, os.path.join(run_dir, INPUTS_DIR)):
        os.makedirs(d, exist_ok=True)
    # 同一个进程跑第二批时必须**换掉**旧的 FileHandler，不能因为"已经有一个"就跳过：
    # 否则第二批的日志会继续写进第一批的目录，新批次目录里连 company_agent.log 都没有。
    # （2026-09-03 排查抓取失败时就被它误导过：14:50 那批的报错写在了 14:30 那批的日志里。）
    for h in [h for h in log.handlers if isinstance(h, logging.FileHandler)]:
        log.removeHandler(h)
        try:
            h.close()
        except Exception:
            pass
    fh = logging.FileHandler(LOG_FILE, encoding="utf-8")
    fh.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s",
                                      datefmt="%H:%M:%S"))
    log.addHandler(fh)
    return run_dir


# ------------------------------------------------------------------------------
# 分类名单（存在分类名单文件时，产品一/二级分类严格按名单映射）
# ------------------------------------------------------------------------------
_DEFAULT_CATEGORY_HINT = (
    "产品分类说明：工作目录下暂无分类名单文件，请按参考词表归类："
    "产品一级分类参考（激光器、光学元件、光电器件、光学仪器、光通信器件、光电材料、"
    "激光设备、光电检测、照明显示、其他）；产品二级分类参考一级分类对应的细分产品。")


def _dedup_pairs(pairs):
    """分类 (一级, 二级) 对去重清洗。"""
    out, seen = [], set()
    for a, b in pairs:
        a = str(a or "").strip().rstrip(";；,，") if a is not None else ""
        b = str(b or "").strip().rstrip(";；,，") if b is not None else ""
        if a.lower() in ("null", "none", "nan"):
            a = ""
        if b.lower() in ("null", "none", "nan"):
            b = ""
        if not a and not b:
            continue
        if (a, b) not in seen:
            seen.add((a, b)); out.append((a, b))
    return out or None


def load_categories():
    """读取分类名单文件（若存在）。返回 ([(一级, 二级), ...], 文件名)；无文件返回 (None, None)。"""
    for path in CATEGORY_FILES:
        if not os.path.exists(path):
            continue
        try:
            if path.lower().endswith(".json"):
                d = json.load(open(path, encoding="utf-8"))
                l1 = d.get("level1") or d.get("一级") or d.get("一级分类") or []
                l2 = d.get("level2") or d.get("二级") or d.get("二级分类") or []
                if l1 and l2:
                    pairs = list(zip(l1, l2))
                elif l1:
                    pairs = [(x, "") for x in l1]
                elif l2:
                    pairs = [("", x) for x in l2]
                else:
                    continue
                pairs = _dedup_pairs(pairs)
                return (pairs, path) if pairs else (None, None)

            if path.lower().endswith(".csv"):
                rows = list(csv.reader(open(path, encoding="utf-8-sig", newline="")))
                if not rows:
                    continue
                header = [c.strip() for c in rows[0]]
                # 表头判定：含“一级/二级/cnName/name”字样则跳过首行
                start = 1 if any(("一级" in h or "二级" in h or h.lower() in ("cnname", "name"))
                                 for h in header) else 0
                i1, i2 = 0, (1 if len(header) >= 2 else None)
                for i, h in enumerate(header):
                    if "一级" in h:
                        i1 = i
                    if "二级" in h:
                        i2 = i
                pairs = []
                for row in rows[start:]:
                    a = row[i1] if i1 is not None and i1 < len(row) else ""
                    b = row[i2] if i2 is not None and i2 < len(row) else ""
                    pairs.append((a, b))
                pairs = _dedup_pairs(pairs)
                return (pairs, path) if pairs else (None, None)

            # xlsx：找“一级/二级”字样列；没有则前两列分别当一/二级，单列当一级
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                data = [r for r in ws.iter_rows(values_only=True)]
                if not data:
                    continue
                header = [str(c).strip() if c is not None else "" for c in data[0]]
                i1 = next((i for i, h in enumerate(header) if "一级" in h), None)
                i2 = next((i for i, h in enumerate(header) if "二级" in h), None)
                if i1 is None and i2 is None:
                    i1 = 0
                    if len(header) >= 2:
                        i2 = 1
                pairs = []
                for row in data[1:]:
                    a = row[i1] if i1 is not None and i1 < len(row) else ""
                    b = row[i2] if i2 is not None and i2 < len(row) else ""
                    pairs.append((a, b))
                pairs = _dedup_pairs(pairs)
                if pairs:
                    wb.close()
                    return pairs, path
            wb.close()
        except Exception as e:
            log.warning("读取分类名单失败 %s: %s", path, e)
    return None, None


def build_categories_hint():
    """生成分类提示语：有分类名单文件则要求严格按名单映射，否则用参考词表。"""
    pairs, path = load_categories()
    if pairs:
        m = OrderedDict()
        for a, b in pairs:
            if a and not b:
                m.setdefault(a, [])
            elif b:
                key = a or "其他"
                m.setdefault(key, [])
                if b not in m[key]:
                    m[key].append(b)
        lines = ["产品分类必须严格从以下分类名单中选取（不要自拟分类名）："]
        if m:
            lines.append("产品一级分类名单（共 %d 个）：%s" % (len(m), "、".join(m.keys())))
            lines.append("各一级分类下属的二级分类名单：")
            for a, bs in m.items():
                lines.append("  %s：%s" % (a, "、".join(bs)))
        lines.append("要求：category_level1 从一级分类名单中取产品最多的前两个（用;分隔）；"
                     "category_level2 必须从对应一级分类的下属二级名单中取前两个（用;分隔）。")
        lines.append("（分类名单来源文件：%s）" % path)
        return "\n".join(lines)
    return _DEFAULT_CATEGORY_HINT
_thread_local = threading.local()
_io_lock = threading.Lock()


def get_session():
    sess = getattr(_thread_local, "session", None)
    if sess is None:
        sess = requests.Session()
        retry = Retry(total=MAX_RETRIES, backoff_factor=1.2,
                      status_forcelist=(429, 500, 502, 503, 504),
                      allowed_methods=frozenset(["GET", "POST"]))
        ad = HTTPAdapter(max_retries=retry, pool_connections=20, pool_maxsize=20)
        sess.mount("http://", ad); sess.mount("https://", ad)
        _thread_local.session = sess
    return sess


# ==============================================================================
# 三、步骤2：多方式抓取 HTML
# ==============================================================================
UA_DESKTOP = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36")
UA_MOBILE = ("Mozilla/5.0 (iPhone; CPU iPhone OS 16_0 like Mac OS X) AppleWebKit/605.1.15 "
             "(KHTML, like Gecko) Version/16.0 Mobile/15E148 Safari/604.1")


def _detect_encoding(resp):
    m = re.search(r"charset=([\w-]+)", resp.headers.get("Content-Type", ""), re.I)
    if m:
        return m.group(1)
    m = re.search(r'charset=["\']?([\w-]+)', resp.content[:3000].decode("ascii", "ignore"), re.I)
    return m.group(1) if m else (resp.apparent_encoding or "utf-8")


def _req_get(url, ua):
    sess = get_session()
    site = "%s://%s" % (urlparse(url).scheme or "http", urlparse(url).netloc)
    resp = sess.get(url, timeout=TIMEOUT, verify=VERIFY_SSL, headers={
        "User-Agent": ua,
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8",
        "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
        "Accept-Encoding": "gzip, deflate",
        "Referer": site + "/",
        "Upgrade-Insecure-Requests": "1",
        "Sec-Fetch-Dest": "document",
        "Sec-Fetch-Mode": "navigate",
        "Sec-Fetch-Site": "same-origin",
        "Sec-Fetch-User": "?1",
    })
    resp.raise_for_status()
    resp.encoding = _detect_encoding(resp)
    return resp.text


def _wanwang_get(url):
    """阿里云万网“网站加速”站点（如 jiepu.com）：正文藏在 CDN 的 Head.js/Body.js 里，
    requests/Playwright 都只能拿到空壳。这里把 JS 载荷解出来重组成 HTML。"""
    sess = get_session()
    site = "%s://%s" % (urlparse(url).scheme or "http", urlparse(url).netloc)
    try:
        resp = sess.get(url, timeout=TIMEOUT, verify=VERIFY_SSL, headers={
            "User-Agent": UA_DESKTOP, "Accept": "text/html,*/*;q=0.8",
            "Accept-Language": "zh-CN,zh;q=0.9", "Referer": site + "/"})
        if resp.status_code != 200:
            return None
        html = resp.text
    except Exception:
        return None
    if "smart-design-mode" not in html and "cdn-static-pages" not in html:
        return None                                   # 非万网加速站点

    def _unescape_js(s):
        s = s.replace("\\r\\n", "\n").replace("\\n", "\n").replace("\\t", "\t")
        s = s.replace("\\'", "'").replace('\\"', '"').replace("\\\\", "\\")
        return re.sub(r"\\u([0-9a-fA-F]{4})", lambda m: chr(int(m.group(1), 16)), s)

    parts = ["<!DOCTYPE html><html><head><meta charset='utf-8'></head><body>"]
    for src in re.findall(r"<script\s+src='([^']+\.(?:Head|Body)\.js[^']*)'", html):
        try:
            js = sess.get(src, timeout=TIMEOUT, verify=VERIFY_SSL,
                          headers={"User-Agent": UA_DESKTOP, "Referer": url}).text
            chunks = re.findall(r"document\.write\(\s*'(.*?)'\s*\)", js, re.S)
            parts.append(_unescape_js("".join(chunks)))
        except Exception:
            continue
    parts.append("</body></html>")
    return "".join(parts)


def _playwright_get(url):
    if not (_HAS_PLAYWRIGHT and USE_PLAYWRIGHT):
        return None
    with sync_playwright() as p:
        b = p.chromium.launch(headless=True)
        # ignore_https_errors: 兼容证书名不匹配的站(ERR_CERT_COMMON_NAME_INVALID)
        ctx = b.new_context(user_agent=UA_DESKTOP, ignore_https_errors=True,
                            locale="zh-CN")
        pg = ctx.new_page()
        try:
            # 用 domcontentloaded 而非 networkidle：万网/wezhan、带长连接的站点
            # 永远等不到 networkidle 会白白超时；DOM 就绪后再固定等一段时间让 JS 渲染
            pg.goto(url, timeout=PW_NAV_TIMEOUT * 1000, wait_until=PW_WAIT_UNTIL)
            pg.wait_for_timeout(PW_RENDER_WAIT)
            html = pg.content()
        finally:
            b.close()
        return html


def _content_len(html):
    """有效正文长度：可见正文字数 + meta description 字数(达到阈值才计入)。
    用于救回服务端已渲染但可见正文很短、简介在 meta 里的站(如资讯/门户站)。"""
    soup = BeautifulSoup(html, "html.parser")
    body_len = len(soup.get_text(strip=True))
    m = soup.find("meta", attrs={"name": "description"})
    desc = (m.get("content") or "").strip() if m else ""
    desc_len = len(desc) if len(desc) >= META_DESC_MIN else 0
    return body_len + desc_len, body_len


def _try_fetch_one(url):
    """对单个 URL 尝试多种方式抓取，返回最佳 HTML 或 None。"""
    methods = [
        ("requests-desktop", lambda: _req_get(url, UA_DESKTOP)),
        ("requests-mobile", lambda: _req_get(url, UA_MOBILE)),
        ("wanwang-CDN解码", lambda: _wanwang_get(url)),
        ("playwright", lambda: _playwright_get(url)),
    ]
    best_html, best_len = None, 0
    for name, fn in methods:
        try:
            html = fn()
            if not html:
                log.warning("抓取方式[%s]无内容 %s", name, url)
                continue
            eff_len, body_len = _content_len(html)
            if eff_len > best_len:
                best_html, best_len = html, eff_len
            if eff_len >= MIN_VALID_HTML:
                log.info("抓取成功[%s] %s", name, url)
                return html
            log.warning("抓取方式[%s]正文过短(%d字, 疑似JS渲染) %s", name, body_len, url)
        except Exception as e:
            log.warning("抓取方式[%s]失败 %s : %s", name, url, e)
    if best_html is not None and best_len >= META_DESC_MIN:
        return best_html
    return None


def _flip_scheme(url):
    """http ↔ https 互切。"""
    if url.startswith("https://"):
        return "http://" + url[8:]
    if url.startswith("http://"):
        return "https://" + url[7:]
    return None


def fetch_html_multi(url):
    """依次尝试多种方式抓 HTML，当前 scheme 全失败则自动切换 http↔https 重试。

    返回 (html, 实际生效的URL)。**必须把生效 URL 一并返回**：
    只返回 HTML 时，协议切换成功后调用方仍以为原 URL 可用，
    最终写进记录 website 的还是那个访问不通的地址（名单里 http/https 写反的很常见）。
    """
    # 第一次尝试：原始 URL
    best = _try_fetch_one(url)
    if best is not None:
        return best, url
    # 如果已有 scheme，切换到另一个重试
    alt = _flip_scheme(url)
    if alt:
        log.info("切换协议重试: %s -> %s", url, alt)
        best = _try_fetch_one(alt)
        if best is not None:
            log.info("协议切换生效，网址更正: %s -> %s", url, alt)
            return best, alt
    log.error("全部方式+协议切换均失败 %s", url)
    return None, url


# ==============================================================================
# 四、工具：清洗 HTML、命名、域名
# ==============================================================================
def base_site(url):
    p = urlparse(url if "://" in url else "http://" + url)
    return "%s://%s" % (p.scheme or "http", p.netloc), p.netloc


def domain_of(netloc):
    host = netloc.split(":")[0].lower()
    return host[4:] if host.startswith("www.") else host


# 二级公共后缀：这些本身不是公司名，公司名在它们**左边**那一段。
# 只列常见的即可——不在表里也只是回退成"剥掉最后一段"，对 xxx.com 这类照样正确。
_PUBLIC_SUFFIX_2 = {
    "com.cn", "net.cn", "org.cn", "gov.cn", "edu.cn", "ac.cn",
    "co.uk", "org.uk", "ac.uk", "gov.uk", "me.uk",
    "co.jp", "ne.jp", "or.jp", "ac.jp", "go.jp",
    "co.kr", "or.kr", "ne.kr",
    "com.tw", "com.hk", "com.mo", "com.sg", "com.my", "com.au", "net.au", "org.au",
    "co.nz", "com.br", "co.in", "com.tr", "co.za", "com.mx", "com.ar",
    "com.vn", "com.ph", "co.th", "com.ru", "com.ua", "co.il",
}


def english_name_from_domain(domain):
    """字段2 英文名称：取域名**注册主体**那一段并首字母大写（fjroe.com -> Fjroe）。

    ⚠ 不能直接 split(".")[0]。很多站点把中文版挂在语言子域上：
      cn.lx-ar.com 会被取成 "Cn"（实测「北京灵犀微光」就是这么错的），
      en. / m. / shop. 之类同理，而 domain_of 只剥掉了 www.。

    这里不用"前缀黑名单"（cn/en/m/zh/shop… 永远列不全），改成
    **从右边剥掉公共后缀，剩下的最后一段就是公司名**，任意层数子域都适用：
        cn.lx-ar.com          -> Lx-ar
        www.ndtek.com         -> Ndtek
        en.abc.example.com.cn -> Example
    """
    labels = [x for x in (domain or "").lower().split(".") if x]
    if not labels:
        return ""
    if len(labels) >= 3 and ".".join(labels[-2:]) in _PUBLIC_SUFFIX_2:
        labels = labels[:-2]          # 剥掉 com.cn / co.uk 这类二级后缀
    elif len(labels) >= 2:
        labels = labels[:-1]          # 剥掉普通顶级后缀
    sld = labels[-1] if labels else ""
    return (sld[:1].upper() + sld[1:]) if sld else ""


def sanitize_filename(name):
    name = re.sub(r'[<>:"/\\|?*\r\n\t]', "_", str(name or "")).strip(" .")
    name = re.sub(r"\s+", " ", name)
    return name[:80] or "未命名"


def clean_html_for_ai(html, keep_links=True):
    """去掉 script/style/注释，压缩空白并截断；keep_links 时保留 <a href> 供找“关于”URL。"""
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "iframe"]):
        tag.decompose()
    if keep_links:
        parts = []
        # 标题与 meta
        if soup.title:
            parts.append("TITLE: " + soup.title.get_text(strip=True))
        for m in soup.find_all("meta"):
            n = (m.get("name") or m.get("property") or "").lower()
            if n in ("description", "keywords", "og:site_name", "og:description"):
                parts.append("META[%s]: %s" % (n, (m.get("content") or "").strip()))
        # 链接（文本 -> href），找“关于”栏目关键
        for a in soup.find_all("a", href=True):
            t = a.get_text(strip=True)
            if t:
                parts.append("LINK: %s -> %s" % (t[:40], a["href"].strip()))
        parts.append("TEXT:\n" + soup.get_text("\n", strip=True))
        text = "\n".join(parts)
    else:
        text = soup.get_text("\n", strip=True)
    return text[:HTML_MAX_CHARS]


# ==============================================================================
# 五、AI 调用（DeepSeek，强制 JSON 输出）
# ==============================================================================
def _salvage_json(text):
    """抢救被截断的 JSON：截到最后一个完整字段，补全未闭合的引号与括号后再解析。"""
    # 去掉最后一个不完整的键值对（最后一个逗号之后的残段）
    candidates = [text]
    last_comma = text.rfind(",")
    if last_comma > 0:
        candidates.append(text[:last_comma])
    for base in candidates:
        s = base
        if s.count('"') % 2:                       # 引号未闭合则补一个
            s += '"'
        s += "]" * (s.count("[") - s.count("]"))   # 补中括号
        s += "}" * (s.count("{") - s.count("}"))   # 补大括号
        try:
            return json.loads(s)
        except Exception:
            continue
    return None


def deepseek_chat(system_prompt, user_prompt):
    """调 DeepSeek 要 JSON。带重试，并处理推理模型特有的「思维链烧穿」。

    ⚠ 2026-09-03 实测踩到的坑：deepseek-v4-flash 是**推理模型**，它先输出思维链
    (reasoning_content) 再输出正文(content)，而**思维链是计入 max_tokens 的**。
    25 字段那个长提示词的思维链要烧 4000~6500 token，一旦超过额度：
        finish_reason=length、completion_tokens 全是 reasoning_tokens、content 为空字符串
    HTTP 还是 200，于是老代码只报一句"模型返回空内容"，看不出到底发生了什么；
    而重试是原样重试，同一个额度撞三次，全废——整家公司就此丢掉。
    实测同一份输入思维链长度在 2836~4096 之间随机波动，所以这是**概率性故障**，
    时好时坏，最难查（北京雅世恒源、北京灵犀微光两家都是栽在这里）。

    ★ 计 token 的规则改不了（思维链计入 max_tokens 是服务端行为），
      但可以**直接把思维链关掉**，那 max_tokens 就只用于正文了。
      本项目全是"照着 HTML 填固定字段"的抽取任务，推理带不来收益：
      实测关掉后同一份输入 5.0s vs 32.2s，10 个关键字段逐一对照**完全一致**。

    对策三条：
      1. 默认 reasoning_effort=none 关掉思维链（AI_REASONING_EFFORT）；
      2. 万一将来该参数失效、又撞上烧穿，第二次重试改用 thinking:{"type":"disabled"}
         这条等效开关（实测同样能把思维链清零），而不是原样再撞一次；
      3. 日志要说人话，把 finish_reason 与思维链 token 数写出来。

    注：曾试过把 model 写成 "deepseek-chat" 来绕开——那不是真实模型（/models 只列出
    v4-flash / v4-pro / v4-flash-vision-exp），接口只是把它当别名路由到 v4-flash 的
    非思考模式。别依赖未文档化的别名，用上面的显式参数。
    """
    cfg = ai_config()               # 网页「设置」页改完立刻生效，无需重启
    headers = {"Authorization": "Bearer " + cfg["api_key"], "Content-Type": "application/json"}
    last_err = None
    no_think = {"reasoning_effort": cfg["reasoning_effort"]} if cfg["reasoning_effort"] else {}
    for attempt in range(MAX_RETRIES):
        payload = {
            "model": cfg["model"],
            "messages": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            "temperature": cfg["temperature"],
            "max_tokens": cfg["max_tokens"],
            "stream": False,
        }
        # 不是所有供应商都支持 response_format，做成开关；关掉时靠提示词要求输出 JSON，
        # 外层本来就有去 ``` 围栏与截断抢救，照样解析得出来
        if cfg["json_mode"]:
            payload["response_format"] = {"type": "json_object"}
        payload.update(no_think)
        try:
            resp = get_session().post(cfg["api_url"], json=payload, headers=headers,
                                      timeout=cfg["timeout"])
            resp.raise_for_status()
            data = resp.json()
            choice = (data.get("choices") or [{}])[0]
            content = (choice.get("message") or {}).get("content") or ""
            content = re.sub(r"^```(?:json)?|```$", "", content.strip(), flags=re.M).strip()
            if not content:
                usage = data.get("usage") or {}
                rt = (usage.get("completion_tokens_details") or {}).get("reasoning_tokens", 0)
                fr = choice.get("finish_reason")
                if fr == "length" and "thinking" not in no_think:
                    log.warning("模型[%s]的思维链吃光了 %d token 额度，正文为空(finish_reason=length)；"
                                "改用 thinking:disabled 重试", cfg["model"], rt)
                    no_think = {"thinking": {"type": "disabled"}}
                    continue                       # 立刻换开关重试，不必等退避
                raise ValueError("模型返回空内容(model=%s, finish_reason=%s, 思维链%d token)"
                                 % (cfg["model"], fr, rt))
            try:
                return json.loads(content)
            except json.JSONDecodeError:
                salvaged = _salvage_json(content)   # 截断时尝试补全未闭合的引号/括号
                if salvaged is not None:
                    log.info("JSON 截断已抢救成功")
                    return salvaged
                raise
        except Exception as e:
            last_err = e
            time.sleep(1.5 * (attempt + 1))
    log.warning("AI 解析失败：%s", last_err)
    return None


# 接口早已不限于 DeepSeek（任何 OpenAI 兼容接口都行），但函数名不能改：
# refill.py / fix_format.py 是用 _ns["deepseek_chat"] 按字符串取的。新代码可用这个别名。
ai_chat = deepseek_chat


# ==============================================================================
# 五·五、电话/传真格式化（按 nation 加国号，统一用 - 分隔）
# ==============================================================================
NATION_CODE = {
    "CN": "86", "HK": "852", "MO": "853", "TW": "886", "JP": "81", "KR": "82",
    "US": "1", "CA": "1", "GB": "44", "UK": "44", "DE": "49", "FR": "33",
    "CH": "41", "IT": "39", "NL": "31", "SE": "46", "SG": "65", "IN": "91",
    "AU": "61", "RU": "7", "ES": "34", "FI": "358", "DK": "45", "BE": "32",
    "AT": "43", "IE": "353", "NO": "47", "PL": "48", "TH": "66", "MY": "60",
    "VN": "84", "ID": "62", "BR": "55", "MX": "52", "IL": "972", "TR": "90",
}
_SEP_RE = re.compile(r"[\s/().．·｜|]+")

# domain 关键字 -> 公司总部国籍（用于校正 nation 字段；按总部口径）
DOMAIN_NATION = {
    "inficon": "CH", "burkert": "DE", "edwardsvacuum": "GB", "edwards": "GB",
    "exfo": "CA", "keyence": "JP", "konicaminolta": "JP", "panasonic": "JP",
    "orion-machinery": "JP", "ulvac": "JP", "osakavacuum": "JP", "usconec": "US",
}


NATION_ALIAS = {
    "CHINA": "CN", "中国": "CN", "UNITEDSTATES": "US", "USA": "US", "AMERICA": "US", "美国": "US",
    "JAPAN": "JP", "日本": "JP", "GERMANY": "DE", "德国": "DE", "UNITEDKINGDOM": "GB", "UK": "GB", "英国": "GB",
    "FRANCE": "FR", "法国": "FR", "SOUTHKOREA": "KR", "KOREA": "KR", "韩国": "KR", "TAIWAN": "TW", "中国台湾": "TW",
    "SWITZERLAND": "CH", "瑞士": "CH", "ITALY": "IT", "意大利": "IT", "CANADA": "CA", "加拿大": "CA",
    "AUSTRALIA": "AU", "澳大利亚": "AU", "RUSSIA": "RU", "俄罗斯": "RU", "SINGAPORE": "SG", "新加坡": "SG",
    "HONGKONG": "HK", "中国香港": "HK", "NETHERLANDS": "NL", "荷兰": "NL", "INDIA": "IN", "印度": "IN",
    "ISRAEL": "IL", "以色列": "IL", "SWEDEN": "SE", "瑞典": "SE", "FINLAND": "FI", "芬兰": "FI",
    "DENMARK": "DK", "丹麦": "DK", "SPAIN": "ES", "西班牙": "ES", "AUSTRIA": "AT", "奥地利": "AT",
    "BELGIUM": "BE", "比利时": "BE", "POLAND": "PL", "波兰": "PL", "THAILAND": "TH", "泰国": "TH",
    "MALAYSIA": "MY", "马来西亚": "MY", "VIETNAM": "VN", "越南": "VN", "INDONESIA": "ID", "印度尼西亚": "ID",
    "BRAZIL": "BR", "巴西": "BR", "MEXICO": "MX", "墨西哥": "MX", "TURKEY": "TR", "土耳其": "TR",
    "NORWAY": "NO", "挪威": "NO", "IRELAND": "IE", "爱尔兰": "IE",
}


def _norm_nation_code(nation):
    """把各种国别写法统一成 ISO 3166-1 两位大写编码；识别不了返回 None。"""
    if not nation:
        return None
    n = str(nation).strip().upper()
    if len(n) == 2 and n.isascii() and n.isalpha():
        return n
    key = re.sub(r"[\s._-]", "", n)
    return NATION_ALIAS.get(key) or NATION_ALIAS.get(n)


def resolve_nation(domain, nation):
    """按 domain + 原 nation 判断公司总部国籍。"""
    d = (domain or "").lower()
    for key, nat in DOMAIN_NATION.items():
        if key in d:
            return nat
    if d.endswith(".jp"):
        return "JP"
    return (nation or "CN").upper()


def _norm_sep(num):
    num = _SEP_RE.sub("-", num.strip())
    return re.sub(r"-{2,}", "-", num).strip("-")


def _format_one(num, code):
    num = num.strip()
    if not num:
        return None
    if num.startswith("+"):                       # 已带国号：仅规整分隔符
        return "+" + _norm_sep(num[1:])
    num = re.sub(r"^00", "", num)                  # 去掉国际接入前缀 00
    groups = [g for g in re.split(r"[^\d]+", num) if g]
    if not groups:
        return None
    if groups[0] == code:                          # 已含国号（0086-xxx 或 86-xxx）：去掉避免重复
        groups = groups[1:]
        if not groups:
            return None
    # 去掉首组的中继前缀 0（意大利除外）：0371 -> 371
    if code != "39" and groups[0].startswith("0"):
        groups[0] = groups[0].lstrip("0") or groups[0]
    return "+%s-%s" % (code, "-".join(groups))


def format_phone(raw, nation):
    """按 nation 加国号并用 - 分隔；支持多号码（逗号/顿号分隔）。无值返回 None。"""
    if raw is None:
        return None
    s = str(raw).strip()
    if not s or s in ("--", "-", "None", "null", "暂无"):
        return None
    code = NATION_CODE.get((nation or "CN").upper(), "86")
    s = re.sub(r"[\(（]\s*0\s*[\)）]", "", s)       # 去掉 (0) 这类中继占位
    out = []
    for p in re.split(r"[,，;；、]|\s/\s", s):
        f = _format_one(p, code)
        if f and f not in out:
            out.append(f)
    return ",".join(out) if out else None


# ==============================================================================
# 六、流程各阶段
# ==============================================================================
def stage2_3_fetch_home(task):
    """步骤2+3：抓首页并落盘 html/<公司名>.html（公司名先用 title 暂定）。"""
    url = task["url"]
    html, eff_url = fetch_html_multi(url)
    if not html:
        task["status"] = "fetch_home_failed"
        return task
    # 协议切换后才抓通 → 以生效网址为准。task["site"] 会被下游用来拼
    # 关于页/联系页 URL、并最终写入 record["website"]，在这里更正即可全链路生效
    if eff_url != url:
        task["url"] = eff_url
        task["site"], _netloc = base_site(eff_url)
        task["url_corrected_from"] = url
    # 暂定公司名（用于文件名）：title 中的公司字样，否则域名
    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else ""
    m = re.search(r"[\u4e00-\u9fa5（）()·]{2,}(?:股份有限公司|有限责任公司|有限公司|集团|研究院|科技|厂|中心)", title)
    provisional = m.group(0) if m else task["domain"]
    task["provisional_name"] = provisional
    path = os.path.join(HTML_DIR, sanitize_filename(provisional) + ".html")
    with _io_lock:
        if os.path.exists(path):
            path = os.path.join(HTML_DIR, sanitize_filename(provisional) + "_" + task["domain"] + ".html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    task["home_html_path"] = path
    task["home_html"] = html
    return task


# ==============================================================================
# 三·五、步骤2b：字段20 log图片 —— 从首页提取并下载 logo
# ==============================================================================
LOGO_MIN_BYTES = 200                    # 小于 200 字节视为无效
LOGO_MAX_BYTES = 2 * 1024 * 1024        # 超过 2MB 视为异常文件
_LOGO_CANDIDATE_GUESSES = ["/favicon.ico", "/logo.png", "/logo.jpg", "/logo.gif",
                           "/images/logo.png", "/images/logo.jpg", "/images/logo.gif",
                           "/static/images/logo.png", "/img/logo.png"]


def _sniff_image_ext(data, content_type, url):
    """按文件魔数/Content-Type/URL后缀判断图片扩展名。"""
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return ".png"
    if data[:3] == b"\xff\xd8\xff":
        return ".jpg"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return ".gif"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return ".webp"
    head = data[:200].lower()
    if b"<svg" in head or head.startswith(b"<?xml"):
        return ".svg"
    ct = (content_type or "").lower()
    for ext in ("png", "jpeg", "gif", "webp", "ico", "svg"):
        if ext in ct:
            return ".jpg" if ext == "jpeg" else "." + ext
    low = url.lower()
    for ext in (".png", ".jpg", ".jpeg", ".gif", ".webp", ".ico", ".svg"):
        if low.endswith(ext):
            return ".jpg" if ext == ".jpeg" else ext
    return ".img"


def _find_logo_candidates(html, site):
    """按优先级找 logo 候选 URL：含 logo 字样的 img -> header 内首图 -> apple-touch-icon -> link icon -> 常见路径。"""
    soup = BeautifulSoup(html, "html.parser")
    cands = []

    def _push(src):
        src = (src or "").strip()
        if src and not src.startswith("data:") and not src.startswith("javascript:"):
            u = urljoin(site + "/", src)
            if is_http_url(u):
                cands.append(u)

    # 1) <img> 的 src/alt/class/id 含 logo 字样
    for img in soup.find_all("img"):
        blob = " ".join(str(x or "") for x in (img.get("src"), img.get("alt"),
                                               img.get("class"), img.get("id"))).lower()
        if "logo" in blob:
            _push(img.get("src"))
    # 2) header 区域内第一张图片（通常就是 logo）
    header = soup.find("header") or soup.find(class_=re.compile(r"\b(head|header|top)\b", re.I))
    if header:
        for img in header.find_all("img"):
            src = img.get("src")
            if src:
                _push(src)
                break
    # 3) apple-touch-icon（手机主屏图标，一般即 logo，分辨率较高）
    for lnk in soup.find_all("link"):
        rel = lnk.get("rel") or []
        rel = rel if isinstance(rel, list) else [rel]
        if any("apple-touch-icon" in str(r).lower() for r in rel):
            _push(lnk.get("href"))
    # 4) favicon 等 link icon
    for lnk in soup.find_all("link"):
        rel = lnk.get("rel") or []
        rel = rel if isinstance(rel, list) else [rel]
        if any("icon" in str(r).lower() for r in rel):
            _push(lnk.get("href"))
    # 5) 常见路径兜底
    for p in _LOGO_CANDIDATE_GUESSES:
        cands.append(site.rstrip("/") + p)
    return cands


def _svg_to_png(svg_bytes):
    """SVG 用 Playwright 渲染截图转 PNG。"""
    if not _HAS_PLAYWRIGHT:
        return None
    import base64
    b64 = base64.b64encode(svg_bytes).decode("ascii")
    with sync_playwright() as p:
        b = p.chromium.launch(headless=True)
        pg = b.new_page()
        try:
            pg.set_content("<html><body style='margin:0'><img id='s' src='data:image/svg+xml;base64,%s'></body></html>" % b64)
            el = pg.query_selector("#s")
            if not el:
                return None
            w = el.evaluate("e => e.naturalWidth") or 0
            h = el.evaluate("e => e.naturalHeight") or 0
            if not w or w > 4000:
                w = 300
            if not h or h > 4000:
                h = 100
            pg.set_viewport_size({"width": max(1, w), "height": max(1, h)})
            return el.screenshot(type="png")
        finally:
            b.close()


def _to_png(data, url=""):
    """把任意格式的 logo 图片统一转成 PNG 字节；失败返回 None。"""
    try:
        from PIL import Image
        import io
        img = Image.open(io.BytesIO(data))
        img.load()
        if getattr(img, "is_animated", False):
            img.seek(0)
        if img.mode in ("RGBA", "LA", "PA") or (img.mode == "P" and "transparency" in img.info):
            img = img.convert("RGBA")
        else:
            img = img.convert("RGB")
        # 超大图缩到 512px 以内，控制文件体积
        if img.width > 512:
            ratio = 512.0 / img.width
            img = img.resize((512, int(img.height * ratio)), Image.LANCZOS)
        buf = io.BytesIO()
        img.save(buf, "PNG")
        return buf.getvalue()
    except Exception:
        pass
    # SVG 等 PIL 不支持的格式：Playwright 渲染截图
    head = data[:600].lstrip().lower()
    if b"<svg" in head or head.startswith(b"<?xml"):
        try:
            return _svg_to_png(data)
        except Exception:
            return None
    return None


def stage_fetch_logo(task):
    """步骤2b：从首页 HTML 找 logo，统一转成 PNG 存到 logos/<域名>.png。"""
    if not task.get("home_html"):
        return task
    seen = []
    for u in _find_logo_candidates(task["home_html"], task["site"]):
        if u not in seen:
            seen.append(u)
    for u in seen[:12]:
        try:
            resp = get_session().get(u, timeout=20, verify=VERIFY_SSL,
                                     headers={"User-Agent": UA_DESKTOP,
                                              "Referer": task["site"] + "/"})
            if resp.status_code != 200:
                continue
            data = resp.content
            if not (LOGO_MIN_BYTES <= len(data) <= LOGO_MAX_BYTES):
                continue
            src_ext = _sniff_image_ext(data, resp.headers.get("Content-Type", ""), u)
            png = _to_png(data, u)                    # 一律统一为 PNG 格式
            if png:
                data, src_ext = png, ".png"
            os.makedirs(LOGO_DIR, exist_ok=True)
            path = os.path.join(LOGO_DIR, task["domain"] + src_ext)
            with _io_lock, open(path, "wb") as f:
                f.write(data)
            task["logo_path"] = path
            log.info("logo 已下载 %s -> %s%s", task["domain"], path,
                     "（已转PNG）" if src_ext == ".png" else "")
            return task
        except Exception as e:
            log.debug("logo 候选失败 %s : %s", u, e)
            continue
    log.warning("未找到可用 logo %s", task["domain"])
    return task


def stage4_parse_home(task):
    """步骤4：AI 解析首页 -> 公司名、产品大类、关于URL、联系我们URL。"""
    if not task.get("home_html"):
        return task
    cleaned = clean_html_for_ai(task["home_html"], keep_links=True)
    result = deepseek_chat(PROMPT_STEP4_SYS, PROMPT_STEP4_USER % cleaned)
    if not result:
        task["ai_name"] = task.get("provisional_name")
        return task
    task["ai_name"] = result.get("cn_name") or task.get("provisional_name")
    task["products"] = result.get("products") or []
    about = result.get("about_url")
    if about and is_http_url(urljoin(task["site"] + "/", str(about).strip())):
        task["about_url"] = urljoin(task["site"] + "/", str(about).strip())
    contact = result.get("contact_url")
    if contact:
        cu = urljoin(task["site"] + "/", str(contact).strip())
        if is_http_url(cu):
            task["contact_url"] = cu
    log.info("首页解析 %s -> %s | 关于URL=%s | 联系URL=%s",
             task["domain"], task["ai_name"], task.get("about_url"), task.get("contact_url"))
    return task


def is_http_url(u):
    """仅 http/https 链接可抓；javascript:/#/mailto: 等一律视为不可抓。"""
    return bool(u) and re.match(r"^https?://", str(u).strip(), re.I) is not None


# 常见"联系我们"页面路径模式（按优先级排列）
_CONTACT_PATH_PATTERNS = [
    "/contact", "/contact-us", "/contactus", "/contact.html", "/contact-us.html",
    "/about/contact", "/about/contact-us", "/about/contactus",
    "/about-us/contact", "/about-us/contact-us",
    "/cn/contact", "/cn/contact-us", "/cn/contactus",
    "/zh/contact", "/zh/contact-us",
    "/lian-xi-wo-men", "/lianxiwomen", "/lxwm", "/lxfs",
    "/contact/index", "/contact/index.html",
]


def _find_contact_links_in_html(html, site):
    """用 BeautifulSoup 从 HTML 中扫描含联系关键词的 <a> 链接。返回候选URL列表。"""
    soup = BeautifulSoup(html, "html.parser")
    candidates = []
    contact_kw = re.compile(r"联系|contact|connect|reach|touch|find\s*us|inquiry|咨询|enquiry",
                            re.I)
    for a in soup.find_all("a", href=True):
        text = a.get_text(strip=True)
        href = a["href"].strip()
        if not href or href.startswith("#"):
            continue
        if contact_kw.search(text):
            cu = urljoin(site + "/", href)
            if is_http_url(cu):
                candidates.append(cu)
    return candidates


def _guess_contact_urls(site):
    """按常见路径模式猜测联系页面URL。"""
    return [site.rstrip("/") + p for p in _CONTACT_PATH_PATTERNS]


def stage5_6_fetch_about(task):
    """步骤5+6：抓“关于”页并落盘 html/关于<公司名>信息.html；无有效关于URL则回退首页。"""
    about_url = task.get("about_url")
    if is_http_url(about_url):
        html, eff = fetch_html_multi(about_url)
        if html and eff != about_url:
            task["about_url"] = eff          # 关于页也可能是协议写反
    else:
        html = None
    if not html:                              # 关于页抓不到/无效则退回首页 HTML
        html = task.get("home_html")
        if not html:
            task["status"] = "fetch_about_failed"
            return task
    name = task.get("ai_name") or task.get("provisional_name") or task["domain"]
    path = os.path.join(HTML_DIR, sanitize_filename("关于" + name + "信息") + ".html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    task["about_html_path"] = path
    task["about_html"] = html
    return task


def stage_find_contact(task):
    """兜底：首页没解析出联系URL时，再从“关于”页 HTML 里解析一次联系URL。"""
    if task.get("contact_url") or not task.get("about_html"):
        return task
    html = task["about_html"]
    site = task["site"]

    # 策略①：AI 解析
    cleaned = clean_html_for_ai(html, keep_links=True)
    result = deepseek_chat(PROMPT_FIND_CONTACT_SYS, PROMPT_FIND_CONTACT_USER % cleaned)
    if result and result.get("contact_url"):
        cu = urljoin(site + "/", str(result["contact_url"]).strip())
        if is_http_url(cu):
            task["contact_url"] = cu
            log.info("AI发现联系URL %s -> %s", task["domain"], cu)
            return task

    # 策略②：BS4 关键词扫描（超链接文本含"联系/contact"等）
    candidates = _find_contact_links_in_html(html, site)
    if candidates:
        task["contact_url"] = candidates[0]
        log.info("关键词扫描发现联系URL %s -> %s", task["domain"], candidates[0])
        return task

    # 策略③：常见路径模式猜测 + 实际验证
    for guess_url in _guess_contact_urls(site):
        html_test, eff = fetch_html_multi(guess_url)
        if html_test is not None:
            task["contact_url"] = eff        # 记生效的那个，而非猜测的原始拼法
            log.info("路径猜测命中联系页 %s -> %s", task["domain"], eff)
            return task

    log.warning("三层策略均未找到联系URL %s", task["domain"])
    return task


def stage_fetch_contact(task):
    """抓“联系我们”页并落盘 html/联系<公司名>信息.html（无有效联系URL则跳过）。"""
    cu = task.get("contact_url")
    if not is_http_url(cu):
        return task
    html, eff = fetch_html_multi(cu)
    if not html:
        log.warning("联系页抓取失败 %s -> %s", task["domain"], cu)
        return task
    if eff != cu:
        task["contact_url"] = eff
    name = task.get("ai_name") or task.get("provisional_name") or task["domain"]
    path = os.path.join(HTML_DIR, sanitize_filename("联系" + name + "信息") + ".html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    task["contact_html_path"] = path
    task["contact_html"] = html
    return task


def stage7_parse_about(task):
    """步骤7：AI 解析关于页 -> 固定 Schema 记录（25 字段）。"""
    if not task.get("about_html"):
        return task
    cleaned = clean_html_for_ai(task["about_html"], keep_links=False)
    products_hint = ";".join(task.get("products") or []) or "（首页未解析出）"
    categories_hint = build_categories_hint()
    result = deepseek_chat(PROMPT_STEP7_SYS,
                           PROMPT_STEP7_USER % (task["url"], products_hint,
                                                categories_hint, cleaned))
    if not result:
        task["status"] = "ai_parse_failed"
        return task
    # 用代码强制覆盖确定性字段，避免模型臆造
    result["website"] = task["site"] + "/"
    result["name"] = english_name_from_domain(task["domain"])
    result.setdefault("cn_name", task.get("ai_name"))
    result["data_source"] = "official_website"
    result["logo"] = task.get("logo_path")           # logo 由程序下载，不采信模型
    result["original_name"] = task.get("original_name")   # 名单里的原始名称，便于比对
    # 仅保留 schema 字段，缺失补 None；_domain 仅用于缓存去重，不导出
    record = {k: result.get(k) for k in SCHEMA_KEYS}
    record["_domain"] = task["domain"]
    task["record"] = record
    task["status"] = "ok"
    return task


CONTACT_FIELDS = ["email", "mobile", "phone", "fax", "address", "cn_address"]


def stage_parse_contact(task):
    """步骤7+：解析“联系我们”页 -> 联系字段，补进主记录（联系页优先填补缺失项）。"""
    if not task.get("contact_html") or not task.get("record"):
        return task
    cleaned = clean_html_for_ai(task["contact_html"], keep_links=False)
    result = deepseek_chat(PROMPT_CONTACT_SYS, PROMPT_CONTACT_USER % cleaned)
    if not result:
        log.warning("联系页解析失败 %s", task["domain"])
        return task
    rec = task["record"]
    filled = []
    for f in CONTACT_FIELDS:
        val = result.get(f)
        if isinstance(val, str):
            val = val.strip() or None
        # 仅补缺失项：关于页该字段为空且联系页有值时才填，避免覆盖已有数据
        if val and not rec.get(f):
            rec[f] = val
            filled.append(f)
    if filled:
        log.info("联系页补全字段 %s -> %s", task["domain"], ",".join(filled))
    task["record"] = rec
    return task


def stage_finalize(task):
    """最后整理：电话按所在地国号格式化，nation 按 domain+nation 校正为总部国籍，并打印完成。"""
    rec = task.get("record")
    if not rec:
        return task
    # 国别先规范为 ISO 两位编码，再按域名校正为公司主体（总部）国籍
    rec["nation"] = resolve_nation(task.get("domain"), _norm_nation_code(rec.get("nation")) or "CN")
    for f in ("phone", "mobile", "fax"):
        rec[f] = format_phone(rec.get(f), rec["nation"])
    fn = _norm_nation_code(rec.get("factory_nation"))
    rec["factory_nation"] = fn if fn else rec["nation"]
    em = re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", str(rec.get("email") or ""))
    rec["email"] = em.group(0) if em else None
    task["record"] = rec
    name = rec.get("cn_name") or rec.get("name") or task.get("domain")
    log.info("%s - 信息整理完毕", name)
    return task


# ==============================================================================
# 七、步骤8：汇总写入（追加合并，不覆盖）
# ==============================================================================
def merge_append(path, new_records, key_func):
    with _io_lock:
        existing = []
        if os.path.exists(path):
            try:
                with open(path, "r", encoding="utf-8") as f:
                    existing = json.load(f)
                if not isinstance(existing, list):
                    existing = []
            except Exception:
                shutil.copy(path, path + ".bak")
                existing = []
        index = {key_func(r): i for i, r in enumerate(existing) if key_func(r)}
        added = updated = 0
        for rec in new_records:
            k = key_func(rec)
            if k and k in index:
                existing[index[k]] = rec; updated += 1
            else:
                if k:
                    index[k] = len(existing)
                existing.append(rec); added += 1
        tmp = path + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(existing, f, ensure_ascii=False, indent=2)
        os.replace(tmp, path)
        return added, updated, len(existing), existing


# ==============================================================================
# 八、步骤9：导出 CSV / Excel
# ==============================================================================
def export_data(records, fmt="excel"):
    """只导出 Excel。CSV 已停用，fmt 仅保留 excel / none 两个取值。"""
    if not records:
        log.info("无数据可导出"); return
    if fmt == "none":
        return
    try:
        from openpyxl import Workbook
    except ImportError:
        log.error("未安装 openpyxl，无法导出 Excel：pip install openpyxl")
        return

    wb = Workbook(); ws = wb.active; ws.title = "公司信息"
    headers = [h for _k, h in SCHEMA_FIELDS]
    ws.append(headers)
    for r in records:
        ws.append([r.get(k) for k in SCHEMA_KEYS])
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, min(50, len(h) * 2))
    ws.freeze_panes = "A2"

    path = os.path.join(EXPORT_DIR, OUTPUT_XLSX)
    tmp = path + ".tmp.xlsx"
    wb.save(tmp)                 # 先写临时文件再替换，
    os.replace(tmp, path)        # 避免导出中途出错把上一次的成品毁掉
    log.info("已导出 Excel -> %s（%d 行）", path, len(records))


# ==============================================================================
# 九、编排
# ==============================================================================
# 待抓名单（按顺序读取并合并去重）。每份表都按列名定位网址列：
# PS_厂家积累表.xlsx 用“官网”列，原始网址名单用“网址”列。
URL_LISTS = [
    "PS_厂家积累表.xlsx",
    os.path.join("原始网址名单", "原始网址_待采集.xlsx"),
    os.path.join("原始网址名单", "剩余名单.xlsx"),
]
# 仅接受形似网址的值（可无 scheme），过滤“暂未找到公开官网”等文本
_URLISH = re.compile(r"^(https?://)?[a-zA-Z0-9][\w.-]*\.[a-zA-Z]{2,}(:\d+)?(/.*)?$")
# 兼容旧的单列格式（第一列网址、第二列清洗状态）
LEGACY_LISTS = ["公司名单_已清洗.xlsx", "公司名单.xlsx"]


def _urls_from_xlsx(path):
    """从一份名单里取 (网址, 原公司名称)。按列名定位：网址列(官网/网址/URL…)，
    公司名称列(公司名称/名称/企业名称/厂家名称)；取不到名称列则名称为空。"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header = [str(c).strip() if c is not None else "" for c in next(rows)]
    except StopIteration:
        wb.close(); return []

    i_url = 0
    for col in ("网址", "官网", "URL", "url", "公司网址"):
        if col in header:
            i_url = header.index(col)
            break
    i_name = None
    for col in ("公司名称", "名称", "企业名称", "厂家名称"):
        if col in header:
            i_name = header.index(col)
            break
    i_status = header.index("清洗状态") if "清洗状态" in header else (1 if len(header) == 2 else None)

    urls = []
    for row in rows:
        if i_url >= len(row):
            continue
        v = str(row[i_url] or "").strip()
        if not v or v.lower() in ("none", "null") or v.startswith("#"):
            continue
        if not _URLISH.match(v):          # 过滤“暂未找到公开官网”等非网址值
            continue
        if i_status is not None and i_status < len(row) \
                and str(row[i_status] or "").strip() == "需人工处理":
            continue
        name = ""
        if i_name is not None and i_name < len(row) and row[i_name] not in (None, ""):
            name = str(row[i_name]).strip()
        urls.append((v, name))
    wb.close()
    return urls


_INPUT_FILES_USED = []          # 本次实际用到的名单文件（load_urls 时填充）


def _copy_input_files():
    """把本次用到的名单文件复制到运行目录/输入名单/ 留档溯源。"""
    dest_dir = os.path.join(RUN_DIR, INPUTS_DIR)
    for path in _INPUT_FILES_USED:
        try:
            if os.path.isfile(path):
                shutil.copy(path, os.path.join(dest_dir, os.path.basename(path)))
                log.info("名单留档: %s -> %s", path, dest_dir)
        except Exception as e:
            log.warning("名单留档失败 %s: %s", path, e)


def load_urls(extra_list=None):
    """优先级: urls.txt > --list 指定名单 > PS_厂家积累表/原始网址名单 > 旧单列名单 > URLS。
    返回 [(url, 原公司名称), ...]，并把用到的名单文件记入 _INPUT_FILES_USED。"""
    global _INPUT_FILES_USED
    _INPUT_FILES_USED = []
    if os.path.exists("urls.txt"):
        urls = []
        with open("urls.txt", "r", encoding="utf-8-sig") as f:
            for line in f:
                line = line.strip()
                if not line or line.startswith("#"):
                    continue
                parts = [p.strip() for p in line.split(",")]
                u = parts[0]
                name = parts[1] if len(parts) > 1 else ""
                if u:
                    urls.append((u, name))
        if urls:
            _INPUT_FILES_USED.append("urls.txt")
            log.info("待抓名单来源: urls.txt（%d 条）", len(urls))
            return urls

    lists = ([extra_list] if extra_list else []) + URL_LISTS
    merged, seen = [], set()
    for path in lists:
        if not path or not os.path.exists(path):
            continue
        try:
            got = _urls_from_xlsx(path)
        except Exception as e:
            log.warning("读取名单失败 %s: %s", path, e)
            continue
        _INPUT_FILES_USED.append(path)
        added = 0
        for u, name in got:
            key = domain_of(urlparse(normalize_url(u)).netloc) or u.lower()
            if key in seen:
                continue
            seen.add(key); merged.append((u, name)); added += 1
        log.info("待抓名单 %s: %d 条，去重后新增 %d", path, len(got), added)
    if merged:
        log.info("待抓队列合计 %d 条（已按域名去重）", len(merged))
        return merged

    for xlsx_path in LEGACY_LISTS:
        if os.path.exists(xlsx_path):
            try:
                urls = _urls_from_xlsx(xlsx_path)
                if urls:
                    _INPUT_FILES_USED.append(xlsx_path)
                    log.info("待抓名单来源(旧格式): %s（%d 条）", xlsx_path, len(urls))
                    return urls
            except Exception:
                pass
    return [(u, "") for u in URLS]


def run_stage(name, fn, items, workers):
    """并发执行一个阶段；单任务异常由 _safe 隔离，不影响同批其它任务。"""
    out = []
    with ThreadPoolExecutor(max_workers=workers) as ex:
        futs = {ex.submit(_safe, fn, t): t for t in items}
        for fu in as_completed(futs):
            out.append(fu.result())
    log.info("阶段[%s]完成", name)
    return out


def _drop_html(tasks, *keys):
    """及时释放已消费的 HTML 文本，降低内存占用（大批量抓取必备）。"""
    for t in tasks:
        for k in keys:
            t.pop(k, None)


def load_processed_domains():
    """已处理域名集合：优先读 JSON 里的 domain，并合并 processed 清单文件。"""
    done = set()
    if os.path.exists(OUTPUT_JSON):
        try:
            with open(OUTPUT_JSON, "r", encoding="utf-8") as f:
                for r in json.load(f):
                    d = r.get("_domain") or r.get("domain")
                    if d:
                        done.add(d)
        except Exception:
            pass
    if os.path.exists(PROCESSED_FILE):
        with open(PROCESSED_FILE, "r", encoding="utf-8") as f:
            for line in f:
                d = line.strip()
                if d:
                    done.add(d)
    return done


def mark_processed(domains):
    """把本批处理过的域名追加进 processed 清单（成功/失败都记，避免重复抓取）。"""
    if not domains:
        return
    with _io_lock, open(PROCESSED_FILE, "a", encoding="utf-8") as f:
        for d in domains:
            f.write(d + "\n")


def process_chunk(chunk_tasks):
    """对一批任务跑完整 8 阶段；阶段间及时释放 HTML 并回收内存。返回成功记录。"""
    t = run_stage("2/3 抓首页", stage2_3_fetch_home, chunk_tasks, FETCH_WORKERS)
    t = run_stage("2b 下载logo", stage_fetch_logo, t, FETCH_WORKERS)
    t = run_stage("4 解析首页(关于+联系URL)", stage4_parse_home, t, AI_WORKERS)
    t = run_stage("5/6 抓关于页", stage5_6_fetch_about, t, FETCH_WORKERS)
    _drop_html(t, "home_html")                       # 关于页已抓完/回退完，首页 HTML 不再需要
    t = run_stage("5b 兜底找联系URL", stage_find_contact, t, AI_WORKERS)
    t = run_stage("6b 抓联系页", stage_fetch_contact, t, FETCH_WORKERS)
    t = run_stage("7 解析关于页", stage7_parse_about, t, AI_WORKERS)
    _drop_html(t, "about_html")
    t = run_stage("7b 解析联系页并合并", stage_parse_contact, t, AI_WORKERS)
    _drop_html(t, "contact_html")
    t = run_stage("8 整理格式化", stage_finalize, t, AI_WORKERS)
    t.sort(key=lambda x: x.get("idx", 0))            # 还原原始顺序
    records = [x["record"] for x in t if x.get("status") == "ok" and x.get("record")]
    failed = [x for x in t if x.get("status") != "ok"]
    return records, failed


def normalize_url(u):
    """确保 URL 有 scheme；www.xxx.com → https://www.xxx.com"""
    u = u.strip()
    if u.startswith("http://") or u.startswith("https://"):
        return u
    return "https://" + u


def run_pipeline(export_fmt, chunk_size=CHUNK_SIZE, force=False, extra_list=None):
    os.makedirs(HTML_DIR, exist_ok=True)
    os.makedirs(CACHE_DIR, exist_ok=True)          # 续跑缓存目录
    urls = load_urls(extra_list)
    _copy_input_files()                            # 名单留档进运行目录/输入名单/

    # 断点续跑：跳过已处理域名，避免重复抓取（4000 家规模的关键）
    done = set() if force else load_processed_domains()
    all_tasks = []
    skipped = 0
    for i, (u, orig_name) in enumerate(urls):
        u = normalize_url(u)        # 确保 URL 始终带 scheme
        site, netloc = base_site(u)
        dom = domain_of(netloc)
        if dom in done:
            skipped += 1
            continue
        all_tasks.append({"idx": i, "url": u, "site": site, "domain": dom,
                          "original_name": orig_name or None, "status": "init"})
    log.info("共 %d 个网址；跳过已处理 %d，待处理 %d，分批大小 %d",
             len(urls), skipped, len(all_tasks), chunk_size)
    if not all_tasks:
        log.info("没有待处理网址，全部已完成")
        if export_fmt and export_fmt != "none":
            export_only(export_fmt)
        return

    total_added = total_updated = total_failed = 0
    n_chunks = (len(all_tasks) + chunk_size - 1) // chunk_size
    for ci in range(n_chunks):
        chunk = all_tasks[ci * chunk_size:(ci + 1) * chunk_size]
        log.info("=== 批次 %d/%d：处理 %d 个网址 ===", ci + 1, n_chunks, len(chunk))
        try:
            records, failed = process_chunk(chunk)
        except Exception:
            # 整批异常（如 OOM）也不让整个程序崩溃：记录后继续下一批
            log.error("批次 %d 整体异常，跳过该批\n%s", ci + 1, traceback.format_exc())
            mark_processed([t["domain"] for t in chunk])
            continue

        # 每批立即落盘 JSON（断点/checkpoint），崩溃最多只丢当前批
        if records:
            a, u, total, _ = merge_append(OUTPUT_JSON, records, lambda r: r.get("_domain") or r.get("domain"))
            total_added += a
            total_updated += u
            log.info("批次 %d 写入：新增 %d / 更新 %d / 累计 %d 条", ci + 1, a, u, total)
        mark_processed([t["domain"] for t in chunk])   # 成功/失败都标记，防重复
        if failed:
            total_failed += len(failed)
            with _io_lock, open(FAILED_FILE, "a", encoding="utf-8") as f:
                for t in failed:
                    f.write("%s\t%s\n" % (t.get("url"), t.get("status")))

        # 主动回收：丢弃本批 tasks 并触发 GC，避免大量 HTML 文本堆积
        chunk.clear()
        gc.collect()

    log.info("全部完成：新增 %d / 更新 %d / 失败 %d", total_added, total_updated, total_failed)

    if export_fmt and export_fmt != "none":
        export_only(export_fmt)


def _safe(fn, task):
    """阶段任务隔离：单任务异常不影响其它任务。"""
    try:
        return fn(task)
    except Exception:
        log.error("阶段任务异常 %s\n%s", task.get("url"), traceback.format_exc())
        task["status"] = "exception"
        return task


def export_only(fmt="excel"):
    """步骤9独立运行：把续跑缓存里的记录导出为 Excel。"""
    if not os.path.exists(OUTPUT_JSON):
        log.error("找不到续跑缓存 %s", OUTPUT_JSON); return
    with open(OUTPUT_JSON, "r", encoding="utf-8") as f:
        records = json.load(f)
    export_data(records, fmt)


def _run_single(url):
    """一键模式：抓取单个 URL，自动续跑到 runs/ 下最新【有数据】的周期目录（无则新建）。"""
    latest = None
    if os.path.isdir(RUNS_DIR):
        dirs = sorted(d for d in os.listdir(RUNS_DIR)
                      if os.path.isdir(os.path.join(RUNS_DIR, d)))
        # 优先选「含续跑缓存数据」的最新周期，避免续跑到空目录导致全量重抓
        for d in reversed(dirs):
            if os.path.exists(os.path.join(RUNS_DIR, d, CACHE_DIR, "companies_info.json")):
                latest = d
                break
        if latest is None:
            latest = dirs[-1] if dirs else None
    run_dir = init_run_env(resume_tag=latest) if latest else init_run_env()
    log.info("一键模式：抓取 %s（续跑周期 %s）", url, run_dir)

    tmp_list = "_single_url.xlsx"
    try:
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        ws.append(["公司名称", "官网"])
        ws.append(["", url])
        wb.save(tmp_list)
    except Exception as e:
        log.error("生成临时名单失败：%s", e)
        return

    try:
        run_pipeline("excel", extra_list=tmp_list)
    finally:
        try:
            os.remove(tmp_list)
        except Exception:
            pass


def main():
    ap = argparse.ArgumentParser(description="LLM 公司信息采集 Agent（周期化运行）")
    ap.add_argument("--export", choices=["excel", "none"], default=DEFAULT_EXPORT,
                    help="抓取完成后是否导出 Excel（默认 excel）")
    ap.add_argument("--export-only", nargs="?", const="excel", choices=["excel"],
                    help="不抓取，仅把续跑缓存导出为 Excel（配合 --resume 使用）")
    ap.add_argument("--chunk-size", type=int, default=CHUNK_SIZE,
                    help="分批大小：每批跑完即落盘并回收内存（默认 %d）" % CHUNK_SIZE)
    ap.add_argument("--force", action="store_true",
                    help="忽略已处理清单，强制重新抓取全部网址")
    ap.add_argument("--list", dest="list_file", default=None, metavar="名单.xlsx",
                    help="指定本次的厂家名单文件（含“官网/网址”列与“公司名称”列）")
    ap.add_argument("--run-name", default=None,
                    help="本次运行批次备注，拼在运行目录名后（如 --run-name 光博会2026夏）")
    ap.add_argument("--resume", default=None, metavar="TAG",
                    help="续跑指定运行目录 runs/<TAG>（断点续跑，不新建目录）")
    args = ap.parse_args()

    # ★ 一键模式：配置了 URL 且未显式传 --list/--resume 时，直接抓取该 URL
    if URL.strip() and not args.list_file and not args.resume:
        _run_single(URL.strip())
        return

    run_dir = init_run_env(args.run_name, args.resume)
    log.info("=" * 70)
    log.info("LLM 公司信息采集 Agent 启动 | 模型=%s | Playwright=%s",
             ai_config()["model"], _HAS_PLAYWRIGHT)
    log.info("运行周期目录: %s", run_dir)
    log.info("=" * 70)

    if args.export_only:
        export_only(args.export_only)
        return
    run_pipeline(args.export, chunk_size=args.chunk_size, force=args.force,
                 extra_list=args.list_file)


if __name__ == "__main__":
    main()
