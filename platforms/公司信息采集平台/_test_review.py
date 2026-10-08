# -*- coding: utf-8 -*-
"""端到端自检：人工审核 + 入库闸门 + 证据服务 + 静态资源。

    python3 _test_review.py

全程不联网、不调 DeepSeek、不连数据库（pymysql 与写库函数都打桩），
审核状态与配置写进临时目录，绝不碰项目里的真实文件。
"""
import json
import os
import shutil
import sys
import tempfile
import types
import urllib.parse

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)

# ---- 打桩 pymysql：WSL 里没装，而且自检本来就不该连真库 ----
fake = types.ModuleType("pymysql")
fake.connect = lambda **kw: (_ for _ in ()).throw(RuntimeError("自检不允许连数据库"))
sys.modules.setdefault("pymysql", fake)

TMP = tempfile.mkdtemp(prefix="审核自检_")
RUN_DIR = os.path.join(BASE, "runs", "__selftest__")
HTML_DIR = os.path.join(RUN_DIR, "html")

from server import review, config                    # noqa: E402
review.STATE_PATH = os.path.join(TMP, "review_state.json")
config.CONFIG_PATH = os.path.join(TMP, "server_config.json")
import script                                        # noqa: E402
script.AI_CONFIG_FILE = config.CONFIG_PATH           # 别让自检读到项目里的真实配置
script._ai_cfg_cache.update(mtime=None, cfg=None)

from fastapi.testclient import TestClient            # noqa: E402
from server import app as appmod, db, mapping        # noqa: E402

client = TestClient(appmod.app)
AUTH = ("admin", "admin123")

PASS = FAIL = 0


def check(label, cond, extra=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print("  ✔ %s" % label)
    else:
        FAIL += 1
        print("  ✘ %s   %s" % (label, extra))


def build_fixture():
    os.makedirs(HTML_DIR, exist_ok=True)
    home = os.path.join(HTML_DIR, "自检公司.html")
    with open(home, "w", encoding="utf-8") as f:
        f.write("<html><head><title>自检</title></head><body>首页正文</body></html>")
    logo_dir = os.path.join(RUN_DIR, "logos")
    os.makedirs(logo_dir, exist_ok=True)
    logo_path = os.path.join(logo_dir, "self.png")
    global LOGO_BYTES
    LOGO_BYTES = bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
        "0000000d4944415478da63fcffff3f0300050001e5ba5a0a0000000049454e44ae426082")
    with open(logo_path, "wb") as f:
        f.write(LOGO_BYTES)

    def rec(cn, site, **kw):
        r = {c: None for c in mapping.DB_COLUMNS}
        r.update({"id": "900000000000001", "cnName": cn, "name": "Selftest Co.",
                  "website": site, "nation": "CN", "location": "Shenzhen",
                  "cnOverview": "<p>一段中文概述</p>", "tags": "激光器;光学元件",
                  "createdAt": "2026-09-03 10:00:00", "updatedAt": "2026-09-03 10:00:00",
                  "updatedBy": "admin", "isEnabled": 1, "check_status": 0})
        r.update(kw)
        return r

    raws = [
        {"cn_name": "自检甲公司", "name": "Selftest A", "website": "https://a.example.com",
         "attribute": "制造商", "category_level1": "激光器", "category_level2": "光纤激光器",
         "nation": "CN", "main_products": "光纤激光器;泵浦源", "company_nature": "激光器"},
        {"cn_name": "自检乙公司", "website": "https://b.example.com",
         "attribute": "代理商", "category_level1": "相机", "category_level2": "工业相机"},
        {"cn_name": "自检丙公司", "website": "https://c.example.com", "category_level1": "光学元件"},
    ]
    mapped = [rec("自检甲公司", "https://a.example.com"),
              rec("自检乙公司", "https://b.example.com"),
              rec("自检丙公司", "https://c.example.com")]
    evid = [
        {"domain": "a.example.com", "site": "https://a.example.com",
         "home_html": os.path.relpath(home, BASE).replace(os.sep, "/"),
         "about_html": None, "contact_html": None,
         "logo": os.path.relpath(logo_path, BASE).replace(os.sep, "/"),
         "about_fallback": True, "contact_missing": True, "products": ["激光器"]},
        {"domain": "b.example.com", "site": "https://b.example.com",
         "about_fallback": False, "contact_missing": False},
        {"domain": "c.example.com", "site": "https://c.example.com"},
    ]
    review.replace_batch(review.build_items(mapped, raws, evid), run_tag="__selftest__")


def main():
    print("临时目录：%s\n" % TMP)
    build_fixture()

    print("【鉴权】")
    check("无凭据被拒", client.get("/api/review/list").status_code == 401)
    check("口令错被拒", client.get("/api/review/list", auth=("admin", "x")).status_code == 401)
    check("正确凭据放行", client.get("/api/review/list", auth=AUTH).status_code == 200)

    print("\n【审核清单】")
    j = client.get("/api/review/list", auth=AUTH).json()
    check("三家公司都在清单里", j["counts"]["total"] == 3, j["counts"])
    check("初始全部未检", j["counts"]["未检"] == 3, j["counts"])
    check("兜底标记透出到列表", j["items"][0]["about_fallback"] is True, j["items"][0])

    print("\n【字段元数据：不能有字段在界面上消失】")
    d = client.get("/api/review/item", params={"key": "a.example.com"}, auth=AUTH).json()
    grouped = set()
    for _g, fs in d["meta"]["groups"]:
        grouped |= set(fs)
    check("34 个 DB 字段全部分了组", grouped == set(mapping.DB_COLUMNS),
          sorted(set(mapping.DB_COLUMNS) - grouped))
    check("25 个原始字段全部可展示", len(d["meta"]["raw_order"]) == 25, len(d["meta"]["raw_order"]))
    check("id/时间戳不可修改",
          not ({"id", "createdAt", "updatedAt"} & set(d["meta"]["editable"])), d["meta"]["editable"])
    check("原始分类字段被标为重点",
          "category_level1" in d["meta"]["raw_key_fields"], d["meta"]["raw_key_fields"])

    print("\n【结论标记】")
    r = client.post("/api/review/mark", json={"key": "a.example.com", "status": "有问题", "note": ""}, auth=AUTH)
    check("「有问题」不写说明被拒", r.status_code == 400, r.text[:120])
    r = client.post("/api/review/mark", json={"key": "a.example.com", "status": "有问题", "note": "一级分类错了"}, auth=AUTH)
    check("带说明可标记有问题", r.status_code == 200 and r.json()["counts"]["有问题"] == 1, r.text[:150])
    r = client.post("/api/review/mark", json={"key": "b.example.com", "status": "通过"}, auth=AUTH)
    check("可标记通过", r.json()["counts"]["通过"] == 1, r.json())
    r = client.post("/api/review/mark", json={"key": "b.example.com", "status": "乱填"}, auth=AUTH)
    check("非法状态被拒", r.status_code == 400)

    print("\n【字段级核对】")
    client.post("/api/review/field", json={"key": "a.example.com", "field": "cnName", "action": "ok"}, auth=AUTH)
    it = client.get("/api/review/item", params={"key": "a.example.com"}, auth=AUTH).json()["item"]
    check("字段可标已核对", it["review"]["fields"].get("cnName") == "ok", it["review"]["fields"])
    r = client.post("/api/review/field", json={"key": "a.example.com", "field": "cnName", "action": "clear"}, auth=AUTH)
    check("可取消字段标记", "cnName" not in r.json()["review"]["fields"], r.json()["review"]["fields"])

    print("\n【人工修正：改动留痕、原值可还原】")
    r = client.post("/api/review/edit", json={"key": "a.example.com", "field": "cnName", "value": "自检甲光电有限公司"}, auth=AUTH)
    check("可修正字段", r.json()["record"]["cnName"] == "自检甲光电有限公司", r.text[:150])
    it = client.get("/api/review/item", params={"key": "a.example.com"}, auth=AUTH).json()["item"]
    check("抓取原值完整保留", it["original"]["cnName"] == "自检甲公司", it["original"]["cnName"])
    check("改动记进 edits", it["review"]["edits"]["cnName"]["from"] == "自检甲公司", it["review"]["edits"])
    check("改过的字段自动算已核对", it["review"]["fields"].get("cnName") == "ok", it["review"]["fields"])
    check("审核历史留痕", any("修正" in h["event"] for h in it["review"]["history"]), it["review"]["history"])
    r = client.post("/api/review/edit", json={"key": "a.example.com", "field": "id", "value": "1"}, auth=AUTH)
    check("系统字段不许改", r.status_code == 400, r.text[:120])
    r = client.post("/api/review/revert", json={"key": "a.example.com", "field": "cnName"}, auth=AUTH)
    check("可还原为抓取原值", r.json()["record"]["cnName"] == "自检甲公司", r.text[:150])
    r = client.post("/api/review/revert", json={"key": "a.example.com", "field": "tags"}, auth=AUTH)
    check("没改过的字段还原被拒", r.status_code == 400)

    print("\n【入库闸门：只写「通过」的】")
    seen = {}
    real_insert = db.insert_records

    def fake_insert(cfg, records, columns=None):
        seen["records"] = records
        return len(records), 0, []
    db.insert_records = fake_insert
    appmod.db.insert_records = fake_insert
    try:
        r = client.post("/api/write-db", auth=AUTH).json()
        names = [x["cnName"] for x in seen.get("records", [])]
        check("只把通过的那 1 家送去写库", names == ["自检乙公司"], names)
        check("提示里说清未写入多少家", "未检" in r["message"] and "有问题" in r["message"], r["message"])
        # 全部退回未检后应拒绝写入
        for k in ("a.example.com", "b.example.com"):
            client.post("/api/review/mark", json={"key": k, "status": "未检"}, auth=AUTH)
        seen.clear()
        r = client.post("/api/write-db", auth=AUTH).json()
        check("没有通过的记录时拒绝写入", r["ok"] is False and "没有" in r["message"], r["message"])
        check("拒绝时根本没调用写库", "records" not in seen)
    finally:
        db.insert_records = real_insert
        appmod.db.insert_records = real_insert

    print("\n【证据服务】")
    r = client.get("/api/evidence", params={"key": "a.example.com", "kind": "home"}, auth=AUTH)
    check("能取到首页原文", r.status_code == 200 and "首页正文" in r.text, r.status_code)
    check("注入了 <base>，页面资源才加载得出来", '<base href="https://a.example.com/">' in r.text)
    r = client.get("/api/evidence", params={"key": "a.example.com", "kind": "about"}, auth=AUTH)
    check("没抓到的页面返回 404 而不是空白", r.status_code == 404, r.status_code)
    check("非法 kind 被拒",
          client.get("/api/evidence", params={"key": "a.example.com", "kind": "x"}, auth=AUTH).status_code == 400)
    check("不存在的公司返回 404",
          client.get("/api/evidence", params={"key": "查无此公司", "kind": "home"}, auth=AUTH).status_code == 404)
    # 路径穿越：证据路径只认审核状态里的值，且必须落在 runs/ 内
    check("runs 之外的路径被挡下", appmod._safe_run_file("server/app.py") is None)
    check("../ 穿越被挡下", appmod._safe_run_file("runs/../server/app.py") is None)
    check("同前缀目录不被放行（runs2 不是 runs）", appmod._safe_run_file("runs2/x.html") is None)

    print("\n【静态资源】")
    for f in ("app.css", "app.js"):
        check("%s 可加载" % f, client.get("/static/" + f, auth=AUTH).status_code == 200)
    check("静态资源禁缓存（改完刷新即生效）",
          client.get("/static/app.js", auth=AUTH).headers.get("cache-control") == "no-store")
    check("静态目录穿越被挡下", client.get("/static/../../server/app.py", auth=AUTH).status_code == 404)
    check("首页可加载", "人工审核" in client.get("/", auth=AUTH).text)

    print("\n【导出：导的是审核后的当前值】")
    client.post("/api/review/mark", json={"key": "c.example.com", "status": "通过"}, auth=AUTH)
    client.post("/api/review/edit", json={"key": "c.example.com", "field": "tags", "value": "人工改过的标签"}, auth=AUTH)
    import io
    from openpyxl import load_workbook

    def sheet(scope):
        r = client.get("/api/export", params={"scope": scope}, auth=AUTH)
        if r.status_code != 200:
            return None, None, r
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb["字段信息"]
        head = [c.value for c in ws[1]]
        rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
        return head, rows, (r, wb)

    head, rows, extra = sheet("all")
    check("导出全部成功", head is not None, extra)
    check("前三列交代审核结论", head[:3] == ["审核状态", "问题说明", "人工修正字段"], head[:4])
    check("sheet1 = 3 审核列 + 默认导出的 26 字段", len(head) == 29, len(head))
    check("★ 默认不导出那 8 个系统维护字段",
          not (set(mapping.EXPORT_DEFAULT_SKIP) & set(head)),
          [h for h in head if h in mapping.EXPORT_DEFAULT_SKIP])
    check("★ 其余字段表头是英文 DB 字段名（人工导库直接对应列名）",
          head[3:] == [c for c in mapping.DB_COLUMNS if c not in mapping.EXPORT_DEFAULT_SKIP],
          head[3:8])
    row_c = next((v for v in rows if v[head.index("cnName")] == "自检丙公司"), None)
    check("人工修正体现在导出里", row_c and row_c[head.index("tags")] == "人工改过的标签", row_c)
    check("审核状态写进了导出", row_c and row_c[0] == "通过", row_c[0] if row_c else None)
    check("修正过哪些字段写进了导出（英文名）", row_c and "tags" in (row_c[2] or ""), row_c[2] if row_c else None)
    check("导出全部含未检的公司",
          any(v[0] == review.UNCHECKED for v in rows), [v[0] for v in rows])
    check("sheet2 是 25 个原始字段", len([c.value for c in extra[1]["原始字段"][1]]) == 25)
    check("sheet2 表头也是英文原始字段名",
          [c.value for c in extra[1]["原始字段"][1]] == [k for k, _h in script.SCHEMA_FIELDS],
          [c.value for c in extra[1]["原始字段"][1]][:5])
    check("sheet2 行数与 sheet1 对齐",
          extra[1]["原始字段"].max_row - 1 == len(rows), extra[1]["原始字段"].max_row - 1)

    head2, rows2, _ = sheet("passed")
    check("导出已通过：只含「通过」的行",
          rows2 and all(v[0] == "通过" for v in rows2), [v[0] for v in (rows2 or [])])
    check("导出已通过的行数少于全部", len(rows2) < len(rows), "%d / %d" % (len(rows2), len(rows)))
    # 把通过的全退回未检，此时"只导通过"应当明确拒绝而不是给一个空表
    for it in review.load()["items"]:
        client.post("/api/review/mark", json={"key": it["key"], "status": "未检"}, auth=AUTH)
    r = client.get("/api/export", params={"scope": "passed"}, auth=AUTH)
    check("没有通过的记录时拒绝导出空表", r.status_code == 400 and "通过" in r.json()["message"], r.text[:120])
    check("此时导出全部仍可用", client.get("/api/export", params={"scope": "all"}, auth=AUTH).status_code == 200)
    client.post("/api/review/mark", json={"key": "c.example.com", "status": "通过"}, auth=AUTH)

    # 按所选公司导出（勾选导出）：只导 keys 里的那几家
    r = client.get("/api/export", params={"keys": "c.example.com"}, auth=AUTH)
    check("按所选公司导出成功", r.status_code == 200, r.status_code)
    wb3 = load_workbook(io.BytesIO(r.content))
    ws3 = wb3["字段信息"]
    head3 = [c.value for c in ws3[1]]
    rows3 = [[c.value for c in row] for row in ws3.iter_rows(min_row=2)]
    check("只导出了所选的那 1 行",
          len(rows3) == 1 and rows3[0][head3.index("cnName")] == "自检丙公司",
          [v[head3.index("cnName")] for v in rows3])
    cd = r.headers.get("content-disposition", "")
    # Content-Disposition 是 percent-encoded 的，先解码再断言（浏览器拿到也是自动解码）
    check("文件名标明了「所选」", "所选" in urllib.parse.unquote(cd), cd)
    # keys 与 scope=passed 叠加：所选里只有通过的会被导
    r = client.get("/api/export", params={"keys": "a.example.com,c.example.com", "scope": "passed"}, auth=AUTH)
    wb4 = load_workbook(io.BytesIO(r.content))
    ws4 = wb4["字段信息"]
    head4 = [c.value for c in ws4[1]]
    rows4 = [[c.value for c in row] for row in ws4.iter_rows(min_row=2)]
    check("所选+已通过组合：只剩通过的那行",
          len(rows4) == 1 and rows4[0][head4.index("cnName")] == "自检丙公司",
          [v[head4.index("cnName")] for v in rows4])
    # 所选公司一个都不在清单里：明确拒绝，别给一张空表
    r = client.get("/api/export", params={"keys": "查无此公司.com"}, auth=AUTH)
    check("所选公司都不存在时拒绝导出", r.status_code == 400 and "所选" in r.json()["message"], r.text[:120])

    # 自选导出字段
    r = client.get("/api/export", params={"fields": "cnName,website"}, auth=AUTH)
    wb5 = load_workbook(io.BytesIO(r.content))
    head5 = [c.value for c in wb5["字段信息"][1]]
    check("自选字段：只导出选中的列", head5[3:] == ["cnName", "website"], head5)
    r = client.get("/api/export", params={"fields": "cnName,updatedAt,type"}, auth=AUTH)
    wb6 = load_workbook(io.BytesIO(r.content))
    head6 = [c.value for c in wb6["字段信息"][1]]
    check("默认排除的字段可以被选回来", head6[3:] == ["cnName", "updatedAt", "type"], head6)
    r = client.get("/api/export", params={"fields": "cnName,不存在的字段"}, auth=AUTH)
    wb7 = load_workbook(io.BytesIO(r.content))
    head7 = [c.value for c in wb7["字段信息"][1]]
    check("不认识的字段名被忽略", head7[3:] == ["cnName"], head7)
    r = client.get("/api/export", params={"fields": "瞎填1,瞎填2"}, auth=AUTH)
    check("一个合法字段都没选时拒绝", r.status_code == 400, r.text[:120])

    print("\n【logo 下载：文件名 = 记录 id】")
    r = client.get("/api/logo-download", params={"key": "a.example.com"}, auth=AUTH)
    check("可下载 logo", r.status_code == 200, r.status_code)
    check("文件名是这家公司的记录 id",
          "900000000000001" in urllib.parse.unquote(r.headers.get("content-disposition", "")),
          r.headers.get("content-disposition", ""))
    check("下载的字节就是文件本身", r.content == LOGO_BYTES, len(r.content))
    check("附件方式下载（attachment）",
          r.headers.get("content-disposition", "").startswith("attachment"), r.headers.get("content-disposition"))
    check("没 logo 的公司返回 404",
          client.get("/api/logo-download", params={"key": "b.example.com"}, auth=AUTH).status_code == 404)
    check("不存在的公司返回 404",
          client.get("/api/logo-download", params={"key": "查无此公司"}, auth=AUTH).status_code == 404)

    r = client.get("/api/export/fields", auth=AUTH).json()
    check("字段清单接口给出全部列与中文名",
          r["columns"] == mapping.DB_COLUMNS and r["cn"]["cnName"] == "中文名称", r["columns"][:4])
    check("字段清单接口给出默认排除的 8 个",
          r["skip"] == sorted(mapping.EXPORT_DEFAULT_SKIP), r["skip"])

    print("\n【AI 接口可配置】")
    r = client.get("/api/ai/presets", auth=AUTH).json()
    ids = [p["id"] for p in r["presets"]]
    check("给出常见供应商预设", len(ids) >= 8 and "deepseek" in ids and "custom" in ids, ids)
    check("每个预设都带完整 chat/completions 地址",
          all(p["api_url"].endswith("/chat/completions") for p in r["presets"] if p["api_url"]),
          [p["api_url"] for p in r["presets"]])

    ai = {"provider": "zhipu", "api_url": "https://open.bigmodel.cn/api/paas/v4/chat/completions",
          "api_key": "test-key-123", "model": "glm-4-flash", "max_tokens": 4096,
          "temperature": 0.2, "timeout": 90, "reasoning_effort": "", "json_mode": False}
    full = {"db": {"host": "127.0.0.1", "port": 3306, "database": "d", "username": "u",
                   "password": "p", "table": "company"},
            "auth": {"username": "admin", "password": "******"},
            "defaults": {"updatedBy": "admin", "isEnabled": 1, "isRecommended": 0,
                         "accessCount": 0, "checkStatus": 0, "type": None, "industryId": 0},
            "ai": ai}
    check("保存 AI 配置", client.post("/api/save-config", json=full, auth=AUTH).json()["ok"] is True)
    pub = client.get("/api/config", auth=AUTH).json()
    check("界面拿回的 Key 已脱敏", pub["ai"]["api_key"] == "******", pub["ai"]["api_key"])
    check("其余 AI 字段原样回显",
          pub["ai"]["model"] == "glm-4-flash" and pub["ai"]["json_mode"] is False, pub["ai"])

    # ★ 换供应商后抓取端必须立刻用新配置（script 读的是同一份文件）
    script._ai_cfg_cache.update(mtime=None, cfg=None)
    eff = script.ai_config()
    check("★ 抓取端立刻用上新配置（不用重启）",
          eff["model"] == "glm-4-flash" and eff["api_key"] == "test-key-123"
          and eff["api_url"].startswith("https://open.bigmodel.cn"), eff)
    check("json_mode 可以关掉", eff["json_mode"] is False, eff["json_mode"])
    check("思维链可置空（不发这个参数）", eff["reasoning_effort"] == "", repr(eff["reasoning_effort"]))
    check("数值字段按类型解析", eff["max_tokens"] == 4096 and eff["timeout"] == 90, eff)

    # 只改数据库、不传 ai 段时，不能把 ai 段冲掉
    full2 = dict(full); full2.pop("ai")
    client.post("/api/save-config", json=full2, auth=AUTH)
    check("不传 ai 段时原样保留（别被数据库那张卡片冲掉）",
          client.get("/api/config", auth=AUTH).json()["ai"]["model"] == "glm-4-flash")
    # 脱敏占位回传时要保留旧 Key
    full3 = dict(full); full3["ai"] = dict(ai, api_key="******", model="glm-4-plus")
    client.post("/api/save-config", json=full3, auth=AUTH)
    script._ai_cfg_cache.update(mtime=None, cfg=None)
    check("回传 ****** 时保留原 Key", script.ai_config()["api_key"] == "test-key-123")
    check("同时改掉了模型名", script.ai_config()["model"] == "glm-4-plus")

    # 测试连接：指向一个必然连不上的地址，验证它如实报错而不是假装成功（不联网）
    bad = dict(ai, api_url="http://127.0.0.1:1/chat/completions", api_key="x")
    r = client.post("/api/test-ai", json=bad, auth=AUTH).json()
    check("测试连接失败时如实报错", r["ok"] is False and "连不上" in r["message"], r["message"])
    r = client.post("/api/ai/models", json=bad, auth=AUTH).json()
    check("拉取模型列表失败时不炸", r["ok"] is False and r["models"] == [], r)

    # 配置坏了不能让抓取起不来
    open(config.CONFIG_PATH, "w", encoding="utf-8").write("{坏文件")
    script._ai_cfg_cache.update(mtime=None, cfg=None)
    check("配置文件损坏时回落默认值",
          script.ai_config()["model"] == script.DEEPSEEK_MODEL, script.ai_config()["model"])
    # 还原一份可用配置，后面的检查还要读它
    with open(config.CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump(full, f, ensure_ascii=False, indent=2)
    script._ai_cfg_cache.update(mtime=None, cfg=None)
    check("修好配置后恢复正常", script.ai_config()["model"] == "glm-4-flash",
          script.ai_config()["model"])

    print("\n【英文名取域名主体】")
    # 曾经直接 split(".")[0]，把 cn.lx-ar.com 取成了 "Cn"（实测「北京灵犀微光」中招）。
    # 现在改成剥掉公共后缀取注册主体；下面这几条锁住行为，别再退回前缀黑名单的做法。
    for dom, want in [
        ("cn.lx-ar.com", "Lx-ar"),          # 中文版挂在语言子域上
        ("www.ndtek.com", "Ndtek"),
        ("ndtek.com", "Ndtek"),
        ("m.example.com", "Example"),
        ("shop.example.com", "Example"),
        ("www.example.com.cn", "Example"),  # 二级公共后缀
        ("cn.example.co.uk", "Example"),
        ("a.b.c.example.co.jp", "Example"), # 多层子域
        ("cn-optics.com", "Cn-optics"),     # ★ 公司名真以 cn- 开头，不能误剥
        ("localhost", "Localhost"),
        ("", ""),
    ]:
        got = script.english_name_from_domain(dom)
        check("%-22s -> %s" % (dom or "(空)", want), got == want, "实际 %r" % got)

    print("\n【字段映射契约】")
    # 打桩翻译，避免自检联网调 DeepSeek
    import server.mapping as _m
    _orig_tr = _m._tr.translate_text
    _m._tr.translate_text = lambda text, lang: "(translated)"
    try:
        row = _m.build_db_record(
            {"cn_name": "自检公司", "name": "", "website": "https://x.example.com",
             "cn_overview": "一段中文简介", "logo": "/tmp/some/logo.png", "tags": "标签"},
            {"updatedBy": "admin", "isEnabled": 1, "isRecommended": 0,
             "accessCount": 0, "checkStatus": 0, "type": None, "industryId": 0})
    finally:
        _m._tr.translate_text = _orig_tr
    check("logoExt 恒为空（logo 要人工传服务器后再填，抓取阶段不许猜后缀）",
          row["logoExt"] is None, row["logoExt"])
    check("即使抓到了 logo 文件 logoExt 也不填", row["logoExt"] is None)
    check("其余字段照常映射", row["cnName"] == "自检公司" and row["website"] == "https://x.example.com", row["cnName"])

    print("\n【状态落盘：重启不丢】")
    check("状态文件已生成", os.path.isfile(review.STATE_PATH))
    saved = json.load(open(review.STATE_PATH, encoding="utf-8"))
    check("落盘内容含记录本身（不只是标记）",
          bool(saved["items"][0]["record"]) and bool(saved["items"][0]["raw"]), list(saved["items"][0]))
    bad = review.STATE_PATH
    with open(bad, "w", encoding="utf-8") as f:
        f.write('{"items": [')
    check("坏文件读为空", review.load()["items"] == [])
    check("坏文件被隔离留证而不是静默清零",
          any(x.startswith("review_state.json.corrupt-") for x in os.listdir(TMP)), os.listdir(TMP))

    print("\n" + "=" * 56)
    print("通过 %d 项，失败 %d 项" % (PASS, FAIL))
    print("=" * 56)
    return FAIL


if __name__ == "__main__":
    try:
        code = main()
    finally:
        shutil.rmtree(TMP, ignore_errors=True)
        shutil.rmtree(RUN_DIR, ignore_errors=True)
    sys.exit(1 if code else 0)
