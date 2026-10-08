# -*- coding: utf-8 -*-
"""统计结果导出：把类别分布统计打成一份可直接给领导看的 Excel。

工作表：
  总览          —— KPI + 数据来源明细
  公司明细      —— 每家公司产品数、占比、覆盖分类数、主力品类
  一级分类分布  —— 含占比与柱状图
  二级分类分布  —— 按一级分组
  公司×一级分类 —— 交叉表（看每家公司的品类结构）
  AI分析        —— 调用任务模型生成的分析报告（可选）

导出范围可只含选中的公司；总量、占比都按所选范围重算。
"""
import os, json, time
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter

from .paths import DATA_DIR
from . import llm as llmmod

EXPORT_DIR = os.path.join(DATA_DIR, "exports")

# 分类色板（浅色底，固定顺序、不循环取色）。与平台前端 --s1/--s2/--s3 同源。
PALETTE = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7", "E34948"]
SEQ = "2A78D6"          # 单系列比大小一律用这一个蓝
MAX_SERIES = 7          # 分类色位上限，超出的尾部折叠成「其他」


def _style_axes(ch):
    """坐标轴与网格保持克制：去掉主图例外的噪声。"""
    ch.y_axis.majorGridlines = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False


def _bar(ws, title, data_ref, cats_ref, anchor, n_cats,
         horizontal=True, color=SEQ, stacked=False, labels=False, x_title="", y_title=""):
    """统一的条形图。单系列不放图例（标题已说明是什么）；多系列必放图例。

    horizontal=True → openpyxl 的 type='bar'（横向条），适合中文长标签；
    type='col' 才是竖向柱。
    """
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.style = 10
    ch.title = title
    if stacked:
        ch.grouping = "stacked"
        ch.overlap = 100
    ch.add_data(data_ref, titles_from_data=True)
    if cats_ref is not None:
        ch.set_categories(cats_ref)
    if x_title:
        ch.x_axis.title = x_title
    if y_title:
        ch.y_axis.title = y_title
    _style_axes(ch)
    n_series = len(ch.series)
    if n_series <= 1:
        ch.legend = None                       # 单系列不需要图例
        if ch.series:
            ch.series[0].graphicalProperties.solidFill = color
            ch.series[0].graphicalProperties.line.solidFill = color
    else:
        ch.legend.position = "b"
        for i, s in enumerate(ch.series):      # 固定顺序取色，绝不循环生成新色
            c = PALETTE[i % len(PALETTE)]
            s.graphicalProperties.solidFill = c
            s.graphicalProperties.line.solidFill = "FCFCFB"   # 段间 2px 表面色间隙
            s.graphicalProperties.line.width = 19050
    if labels:
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        ch.dataLabels.showSerName = False
        ch.dataLabels.showCatName = False
    # 图表数据源放在隐藏的辅助列里，必须关掉"只绘制可见单元格"，否则图是空的
    # （openpyxl 里叫 visible_cells_only，对应 Excel 的 plotVisOnly）
    ch.visible_cells_only = False
    ch.height = max(7.0, min(0.62 * n_cats + 3.2, 26))
    ch.width = 22 if not stacked else 26
    ws.add_chart(ch, anchor)
    return ch

H_FILL = PatternFill("solid", fgColor="2A78D6")
H_FONT = Font(color="FFFFFF", bold=True, size=11)
T_FONT = Font(bold=True, size=14)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


ANALYZE_PROMPT = """你是一名面向公司管理层的产品数据分析师。下面是本平台统计出的产品数据分布（JSON）。
请据此写一份**给领导看**的分析报告。硬性要求：

1. **结论先行**：每一节开头先用一句话给结论，再用数据支撑。
2. **只用给定数据**，绝对不得编造任何数字、公司名或事实；数据里没有的，就说"当前数据未覆盖"。
3. 面向决策，语言简洁；不要罗列原始明细，不要谈技术实现和数据处理流程。
4. 用 Markdown 小标题分节，全文中文，控制在 900 字以内。

## 一、总体概况
产品总量、覆盖公司数、覆盖的一级/二级分类数。一句话概括当前产品数据资产的规模与完整度。

## 二、类别分布与集中度
产品主要集中在哪几个一级分类（给出具体数量和占比）？前三大类合计占比说明了什么？
哪些一级分类覆盖薄弱甚至空白？对产品线布局意味着什么。

## 三、公司梯队
按产品数量把公司分为头部/腰部/长尾三档，指出头部公司及其占比。
是否存在单一公司占比过高、整体结构失衡的情况，这带来什么风险。

## 四、优势细分领域
在产品量最大的两到三个一级分类里，具体集中在哪些二级分类？反映我们在哪些细分领域积累最深。

## 五、给管理层的建议
3-5 条可执行建议，按优先级排列。例如：哪些薄弱品类值得优先补充采集、
哪些公司的数据值得深挖、分类颗粒度是否存在问题。每条建议后用括号注明支撑数据。

统计数据：
"""


def _sheet(wb, title, headers, rows, widths=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = H_FILL, H_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = BORDER
    return ws


def _pct(n, total):
    return round(n * 100.0 / total, 2) if total else 0.0


def build_workbook(dist, companies, analysis="", note=""):
    """dist = stats.scan_distribution 的结果；companies = 选中的公司名列表（空=全部）。"""
    sel = [c for c in dist["companies"] if not companies or c in companies]
    if not sel:
        raise ValueError("所选范围内没有任何公司数据")
    by = dist["by_company"]
    total = sum(by[c]["total"] for c in sel)

    cat1 = Counter()
    cat2 = Counter()
    cat1_comp = {}                      # 一级分类 → 涉及公司数
    for c in sel:
        for x in by[c]["cat1"]:
            cat1[x["name"]] += x["n"]
            cat1_comp.setdefault(x["name"], set()).add(c)
        for x in by[c]["cat2"]:
            cat2[x["name"]] += x["n"]

    wb = Workbook()

    # ---------- 总览 ----------
    ws = wb.active
    ws.title = "总览"
    ws["A1"] = "产品数据分布统计报告"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"生成时间：{time.strftime('%Y-%m-%d %H:%M:%S')}　|　数据扫描时间：{dist.get('scanned_at', '')}"
    ws["A2"].font = Font(size=10, color="808080")
    ws["A3"] = ("统计范围：全部公司" if not companies else f"统计范围：所选 {len(sel)} 家公司")
    ws["A3"].font = Font(size=10, color="808080")
    if note:
        ws["A4"] = note
        ws["A4"].font = Font(size=10, color="808080")

    kpi = [("产品总数", total), ("覆盖公司数", len(sel)),
           ("覆盖一级分类数", len(cat1)), ("覆盖二级分类数", len(cat2)),
           ("平均每家公司产品数", round(total / len(sel), 1) if sel else 0)]
    r0 = 6
    ws.cell(row=r0, column=1, value="核心指标").font = T_FONT
    for i, (k, v) in enumerate(kpi):
        ws.cell(row=r0 + 1 + i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r0 + 1 + i, column=2, value=v)
    r1 = r0 + len(kpi) + 3
    ws.cell(row=r1, column=1, value="数据来源").font = T_FONT
    ws.cell(row=r1 + 1, column=1, value="来源").font = H_FONT
    ws.cell(row=r1 + 1, column=1).fill = H_FILL
    for j, h in enumerate(["公司数", "产品数"], start=2):
        c = ws.cell(row=r1 + 1, column=j, value=h)
        c.font, c.fill = H_FONT, H_FILL
    srcs = [s for s in dist.get("sources", []) if s["products"]]
    for i, s in enumerate(srcs):
        ws.cell(row=r1 + 2 + i, column=1, value=s["label"])
        ws.cell(row=r1 + 2 + i, column=2, value=s["companies"])
        ws.cell(row=r1 + 2 + i, column=3, value=s["products"])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14

    # 图：来源构成 = 部分-整体 → 单条横向堆叠条（不用饼图）
    # 堆叠要求每个来源是独立系列，所以把数据横排到辅助区（H 列起，随后隐藏）
    if srcs:
        hr = r1 + 1
        for j, s in enumerate(srcs):
            ws.cell(row=hr, column=8 + j, value=s["label"])
            ws.cell(row=hr + 1, column=8 + j, value=s["products"])
        ws.cell(row=hr + 1, column=7, value="全部产品")
        data = Reference(ws, min_col=8, max_col=8 + len(srcs) - 1, min_row=hr, max_row=hr + 1)
        cats = Reference(ws, min_col=7, min_row=hr + 1, max_row=hr + 1)
        ch = _bar(ws, f"产品来源构成（合计 {total} 个产品）", data, cats, "E6", 1,
                  horizontal=True, stacked=True, labels=True)
        ch.height, ch.width = 6.5, 24
        for col in range(7, 8 + len(srcs)):
            ws.column_dimensions[get_column_letter(col)].hidden = True

    # ---------- 公司明细 ----------
    rows = []
    for c in sorted(sel, key=lambda x: -by[x]["total"]):
        e = by[c]
        top = "、".join(f"{x['name']}({x['n']})" for x in e["cat1"][:3])
        rows.append([c, e.get("source", ""), e["total"], _pct(e["total"], total),
                     len(e["cat1"]), len(e["cat2"]), top])
    ws1 = _sheet(wb, "公司明细",
                 ["公司", "数据来源", "产品数", "占比%", "覆盖一级分类", "覆盖二级分类", "主力品类 Top3"],
                 rows, [34, 16, 10, 10, 14, 14, 52])
    # 图：33 家排名比大小 → 横向条形图（单系列单色）。表已按产品数降序，取前 15 家。
    if rows:
        n = min(15, len(rows))
        data = Reference(ws1, min_col=3, min_row=1, max_row=n + 1)
        cats = Reference(ws1, min_col=1, min_row=2, max_row=n + 1)
        _bar(ws1, f"公司产品数 Top {n}（共 {len(rows)} 家，完整名单见左表）",
             data, cats, "I2", n, labels=True)

    # ---------- 一级分类分布 ----------
    rows = [[k, n, _pct(n, total), len(cat1_comp.get(k, ()))] for k, n in cat1.most_common()]
    ws2 = _sheet(wb, "一级分类分布", ["一级分类", "产品数", "占比%", "涉及公司数"],
                 rows, [26, 12, 12, 14])
    # 图：12 个类比大小 → 横向条形图，中文标签横排更好读
    if rows:
        data = Reference(ws2, min_col=2, min_row=1, max_row=len(rows) + 1)
        cats = Reference(ws2, min_col=1, min_row=2, max_row=len(rows) + 1)
        _bar(ws2, "各一级分类产品数", data, cats, "F2", len(rows), labels=True)

    # ---------- 二级分类分布 ----------
    rows = []
    for key, n in cat2.most_common():
        c1, _, c2 = key.partition("/")
        rows.append([c1, c2, n, _pct(n, total)])
    ws3 = _sheet(wb, "二级分类分布", ["一级分类", "二级分类", "产品数", "占比%"],
                 rows, [24, 30, 12, 12])
    # 图：86 个二级分类全画上去没法看 → 只画 Top 20（表里是全量）
    if rows:
        n = min(20, len(rows))
        data = Reference(ws3, min_col=3, min_row=1, max_row=n + 1)
        cats = Reference(ws3, min_col=2, min_row=2, max_row=n + 1)
        _bar(ws3, f"二级分类产品数 Top {n}（共 {len(rows)} 个，完整清单见左表）",
             data, cats, "F2", n, labels=True)

    # ---------- 公司 × 一级分类 交叉表 ----------
    order = [k for k, _ in cat1.most_common()]
    rows = []
    for c in sorted(sel, key=lambda x: -by[x]["total"]):
        m = {x["name"]: x["n"] for x in by[c]["cat1"]}
        rows.append([c] + [m.get(k, 0) for k in order] + [by[c]["total"]])
    rows.append(["合计"] + [cat1[k] for k in order] + [total])
    ws4 = _sheet(wb, "公司×一级分类", ["公司"] + order + ["合计"], rows,
                 [34] + [12] * len(order) + [10])
    # 图：每家公司的品类构成（部分-整体）→ 横向堆叠条。
    # 12 个一级分类做系列会超过 7 个色位上限，取前 6 大类 + 其余合并为「其他」。
    if len(rows) > 1:
        keep = order[:MAX_SERIES - 1]
        rest = order[MAX_SERIES - 1:]
        top_comp = [r for r in rows[:-1]][:15]
        hc0 = len(order) + 3                       # 辅助区起始列（表格右侧，随后隐藏）
        ws4.cell(row=1, column=hc0, value="公司")
        for j, k in enumerate(keep):
            ws4.cell(row=1, column=hc0 + 1 + j, value=k)
        if rest:
            ws4.cell(row=1, column=hc0 + 1 + len(keep), value=f"其他{len(rest)}类")
        for i, r in enumerate(top_comp):
            ws4.cell(row=2 + i, column=hc0, value=r[0])
            m = dict(zip(order, r[1:1 + len(order)]))
            for j, k in enumerate(keep):
                ws4.cell(row=2 + i, column=hc0 + 1 + j, value=m.get(k, 0))
            if rest:
                ws4.cell(row=2 + i, column=hc0 + 1 + len(keep),
                         value=sum(m.get(k, 0) for k in rest))
        ncol = len(keep) + (1 if rest else 0)
        data = Reference(ws4, min_col=hc0 + 1, max_col=hc0 + ncol,
                         min_row=1, max_row=1 + len(top_comp))
        cats = Reference(ws4, min_col=hc0, min_row=2, max_row=1 + len(top_comp))
        _bar(ws4, f"各公司品类构成 Top {len(top_comp)}（按产品数，前 {len(keep)} 大类单列、其余合并）",
             data, cats, f"A{len(rows) + 4}", len(top_comp), stacked=True)
        for col in range(hc0, hc0 + ncol + 1):
            ws4.column_dimensions[get_column_letter(col)].hidden = True

    # ---------- AI 分析 ----------
    if analysis:
        ws5 = wb.create_sheet("AI分析")
        ws5.column_dimensions["A"].width = 118
        ws5["A1"] = "AI 分析报告"
        ws5["A1"].font = Font(bold=True, size=14)
        r = 3
        for line in analysis.split("\n"):
            t = line.rstrip()
            cell = ws5.cell(row=r, column=1, value=t.lstrip("#").strip() if t.startswith("#") else t)
            if t.startswith("#"):
                cell.font = Font(bold=True, size=12, color="2A78D6")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
    return wb


def _taxonomy_size():
    """标准分类体系规模。必须喂给模型 —— 否则它会自己编一个总数来算覆盖率
    （实测编成过 162，真实是 160）。"""
    try:
        from .paths import CONFIG_DIR
        cats = json.load(open(os.path.join(CONFIG_DIR, "company_categories.json"),
                              encoding="utf-8"))["categories"]
        return len(cats), sum(len(v["children"]) for v in cats.values())
    except Exception:
        return 0, 0


def _payload_for_ai(dist, companies):
    sel = [c for c in dist["companies"] if not companies or c in companies]
    by = dist["by_company"]
    total = sum(by[c]["total"] for c in sel)
    cat1, cat2 = Counter(), Counter()
    for c in sel:
        for x in by[c]["cat1"]:
            cat1[x["name"]] += x["n"]
        for x in by[c]["cat2"]:
            cat2[x["name"]] += x["n"]
    n1, n2 = _taxonomy_size()
    uncovered = []
    if n1:
        try:
            from .paths import CONFIG_DIR
            allc = json.load(open(os.path.join(CONFIG_DIR, "company_categories.json"),
                                  encoding="utf-8"))["categories"]
            uncovered = [k for k in allc if k not in cat1]
        except Exception:
            uncovered = []
    return {
        "产品总数": total, "公司数": len(sel),
        "标准分类体系": {"一级分类总数": n1, "二级分类总数": n2,
                       "说明": "覆盖率请以此为分母计算，不要使用其他数字"},
        "覆盖一级分类数": len(cat1), "覆盖二级分类数": len(cat2),
        "完全没有产品的一级分类": uncovered,
        "一级分类分布": [{"分类": k, "产品数": n, "占比%": _pct(n, total)}
                        for k, n in cat1.most_common()],
        "二级分类分布Top30": [{"分类": k, "产品数": n, "占比%": _pct(n, total)}
                            for k, n in cat2.most_common(30)],
        "各公司产品数": [{"公司": c, "产品数": by[c]["total"], "占比%": _pct(by[c]["total"], total),
                        "主力品类": [x["name"] for x in by[c]["cat1"][:3]]}
                       for c in sorted(sel, key=lambda x: -by[x]["total"])],
    }


def run_export(job, dist, companies, model_cfg=None, note=""):
    """后台任务：可选调用 LLM 生成分析 → 生成 xlsx → 返回文件路径。"""
    os.makedirs(EXPORT_DIR, exist_ok=True)
    sel_n = len([c for c in dist["companies"] if not companies or c in companies])
    job.step(0, 3)
    job.log(f"统计范围：{'全部' if not companies else f'所选 {sel_n} 家'}公司")

    analysis = ""
    if model_cfg:
        job.set_current("正在调用任务模型生成分析报告 …")
        job.log("调用 AI 生成面向管理层的分析报告（可能需要 10-60 秒）…")
        payload = _payload_for_ai(dist, companies)
        text, err = llmmod.chat(model_cfg,
                                ANALYZE_PROMPT + json.dumps(payload, ensure_ascii=False, indent=1))
        if err:
            job.log(f"⚠ AI 分析调用失败，将只导出统计表：{err}")
        else:
            analysis = text or ""
            job.log(f"AI 分析完成（{len(analysis)} 字）")
    else:
        job.log("未选择任务模型，本次只导出统计表（不含 AI 分析）")
    job.step(1)

    job.set_current("正在生成 Excel …")
    wb = build_workbook(dist, companies, analysis, note)
    job.step(2)

    name = f"产品数据分布统计_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    path = os.path.join(EXPORT_DIR, f"{job.id}.xlsx")
    wb.save(path)
    job.step(3)
    job.log(f"✓ 导出完成：{name}（{os.path.getsize(path) // 1024} KB）")
    return {"file": path, "name": name, "companies": sel_n,
            "has_analysis": bool(analysis), "analysis": analysis,
            "size": os.path.getsize(path)}
